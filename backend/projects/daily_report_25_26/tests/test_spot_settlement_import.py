"""原始账单回归；真实数据库验证使用外层事务并最终回滚。"""
import base64
import io
import json
import os
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event, text
from sqlalchemy.orm import Session

from backend.config import DATA_DIRECTORY
from backend.db.database_daily_report_25_26 import engine
from backend.projects.daily_report_25_26.services import spot_settlement_import as service
from backend.projects.daily_report_25_26.api import spot_settlement as api

FIXTURES = None
NAMES = ['大连洁净能源集团有限公司2026-01-12.xlsx', '大连热电股份有限公司北海热电厂2026-01-12.xlsx']


def fixtures():
    if FIXTURES is not None:
        return FIXTURES
    folder = Path(__file__).resolve().parents[4] / 'configs'
    return {name: (folder / name).read_bytes() for name in NAMES}


def edited(content, callback):
    book = load_workbook(io.BytesIO(content))
    callback(book)
    result = io.BytesIO()
    book.save(result)
    book.close()
    return result.getvalue()


class BillParserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.files = fixtures()
        cls.reports = [service.parse_bill(cls.files[name], name) for name in NAMES]

    def test_all_units_and_all_subjects_including_blanks(self):
        self.assertEqual([len(p['units']) for p in self.reports], [4, 3])
        self.assertEqual([len(p['records']) for p in self.reports], [4268, 3201])
        for report in self.reports:
            for unit in report['units']:
                rows = [r for r in report['records'] if r['settlement_unit_id'] == unit['settlement_unit_id']]
                self.assertEqual(len(rows), 97 * 11)
                self.assertEqual({r['point_index'] for r in rows}, set(range(97)))
                self.assertTrue(any(r['quantity'] is None and r['price'] is None and r['fee'] is None for r in rows))

    def test_zero_unit_negative_values_and_original_coordinates(self):
        report = self.reports[0]
        unit = next(u for u in report['units'] if u['unit_number'] == 2)
        self.assertEqual(unit['unit_metered_quantity'], 0)
        rows = [r for r in report['records'] if r['settlement_unit_id'] == unit['settlement_unit_id']]
        self.assertEqual(len(rows), 1067)
        first = next(r for r in report['records'] if r['unit_number'] == 1 and r['point_index'] == 1 and r['subject_code'] == '100600')
        self.assertEqual(str(first['fee']), '-554.4')
        self.assertEqual(first['price'], -100)
        self.assertEqual(first['source_cells']['fee']['cell'], 'G13')
        self.assertEqual(first['source_cells']['fee']['value'], -554.4)

    def test_final_point_and_daily_totals_are_separate(self):
        for report in self.reports:
            last = next(r for r in report['records'] if r['point_index'] == 96)
            self.assertEqual(last['time_label'], '24:00')
            self.assertEqual(last['settlement_date'].isoformat(), '2026-01-12')
            daily = next(r for r in report['records'] if r['record_type'] == 'daily_total')
            self.assertIsNone(daily['price'])
            self.assertNotIn('price', daily['source_cells'])
            # 日合计不再与分时明细相加，保持两个原始层级。
            summed = sum(r['fee'] or 0 for r in report['records'] if r['record_type'] == 'period')
            self.assertEqual(summed, report['units'][0]['bill_total_fee'])

    def test_duplicate_missing_and_invalid_period_reject_entire_file(self):
        callbacks = [lambda b: setattr(b.active['A14'], 'value', '00:15'),
                     lambda b: b.active.delete_rows(11, 3),
                     lambda b: setattr(b.active['A11'], 'value', '24:15')]
        for callback in callbacks:
            with self.assertRaises(ValueError):
                service.parse_bill(edited(self.files[NAMES[0]], callback), NAMES[0])

    def test_unknown_sheet_and_conflicting_dates_not_silently_ignored(self):
        def extra(book):
            book.create_sheet('未识别')['A1'] = '其他内容'
        for callback in [extra, lambda b: setattr(b.active['A1'], 'value', '2026-01-13日账单')]:
            with self.assertRaises(ValueError):
                service.parse_bill(edited(self.files[NAMES[0]], callback), NAMES[0])

    def test_formula_and_text_preserved_without_inventing_numeric_values(self):
        def change(book):
            book.active['F11'] = '—'
            book.active['F12'] = '=1+1'
        report = service.parse_bill(edited(self.files[NAMES[0]], change), NAMES[0])
        row = next(r for r in report['records'] if r['unit_number'] == 1 and r['point_index'] == 1 and r['subject_code'] == '100401')
        self.assertIsNone(row['quantity'])
        self.assertIsNone(row['price'])
        self.assertEqual(row['source_cells']['quantity']['value'], '—')
        self.assertEqual(row['source_cells']['price']['value'], '=1+1')
        self.assertEqual(len(report['warnings']), 2)

    def test_preview_deduplicates_and_owner_scope_checked_before_database(self):
        with tempfile.TemporaryDirectory(dir=DATA_DIRECTORY) as folder:
            preview = service.stage_uploads([(NAMES[0], self.files[NAMES[0]])] * 2, 'tester', folder)
            self.assertEqual(len(preview['files']), 1)
            with self.assertRaises(PermissionError):
                service.persist_upload(preview['token'], 'other', root=folder)
            with self.assertRaises(PermissionError):
                service.persist_upload(preview['token'], 'tester', allowed_units={'BeiHai'}, root=folder)

    def test_api_preview_confirm_and_invalid_token(self):
        app = FastAPI()
        app.include_router(api.router)
        app.dependency_overrides[api._importer] = lambda: SimpleNamespace(username='tester', group='Global_admin')
        with tempfile.TemporaryDirectory(dir=DATA_DIRECTORY) as folder, patch.object(service, 'IMPORT_ROOT', Path(folder)):
            client = TestClient(app)
            reply = client.post('/spot_electricity/bills/preview', files=[('files', (NAMES[0], self.files[NAMES[0]], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))])
            self.assertEqual(reply.status_code, 200, reply.text)
            token = reply.json()['token']
            with patch.object(api, 'persist_upload', return_value={'ok': True, 'inserted': 4268}) as save:
                self.assertTrue(client.post('/spot_electricity/bills/confirm', json={'token': token}).json()['ok'])
                self.assertEqual(save.call_args.args[1], 'tester')
            self.assertEqual(client.post('/spot_electricity/bills/confirm', json={'token': '../../escape'}).status_code, 422)

    def test_actual_session_group_shape_and_unit_scope(self):
        # 真实 AuthSession 使用 group 字段，没有 is_global_admin 属性。
        admin = SimpleNamespace(group='Global_admin', username='tester')
        self.assertIs(api._importer(admin), admin)
        self.assertIsNone(api._allowed(admin))
        unit = SimpleNamespace(group='unit_admin', has_project_access=lambda _: True,
                               get_project_action_flags=lambda _: SimpleNamespace(can_submit=True),
                               resolve_allowed_units=lambda _: {'BeiHai'})
        self.assertEqual(api._allowed(unit), {'BeiHai'})
        with self.assertRaises(Exception) as raised:
            api._importer(unit)
        self.assertEqual(raised.exception.status_code, 403)


class BillAccessTests(unittest.TestCase):
    def setUp(self):
        app = FastAPI()
        app.include_router(api.router)
        self.app = app
        self.client = TestClient(app)

    def test_project_submitter_cannot_use_any_bill_endpoint(self):
        self.app.dependency_overrides[api.get_current_session] = lambda: SimpleNamespace(
            username='submitter', group='unit_admin', has_project_access=lambda _: True,
            get_project_action_flags=lambda _: SimpleNamespace(can_submit=True))
        requests = [
            ('post', '/preview', {'files': [('files', ('test.xlsx', b'test'))]}),
            ('post', '/confirm', {'json': {'token': '00000000-0000-0000-0000-000000000001'}}),
            ('get', '/imports', {}),
            ('get', '/records', {'params': {'import_id': '00000000-0000-0000-0000-000000000001'}}),
        ]
        for method, path, kwargs in requests:
            response = getattr(self.client, method)('/spot_electricity/bills' + path, **kwargs)
            self.assertEqual(response.status_code, 403, response.text)

    def test_admin_can_preview_more_than_twenty_files(self):
        self.app.dependency_overrides[api.get_current_session] = lambda: SimpleNamespace(username='admin', group='Global_admin')
        files = [('files', (f'{i}.xlsx', b'fixture')) for i in range(25)]
        with patch.object(api, 'stage_uploads', return_value={'ok': True, 'files': [], 'token': 'preview'}) as stage:
            response = self.client.post('/spot_electricity/bills/preview', files=files)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(len(stage.call_args.args[0]), 25)


class BillDatabaseTests(unittest.TestCase):
    def setUp(self):
        if os.environ.get('SPOT_BILL_TEST_DB') != '1':
            self.skipTest('需显式启用真实数据库回滚测试')
        self.connection = engine.connect()
        self.outer = self.connection.begin()
        self.factory = lambda: Session(bind=self.connection, join_transaction_mode='create_savepoint')
        self.directory = tempfile.TemporaryDirectory(dir=DATA_DIRECTORY)
        self.root = Path(self.directory.name)
        self.files = fixtures()

    def tearDown(self):
        if hasattr(self, 'outer'):
            self.outer.rollback()
            self.connection.close()
            self.directory.cleanup()

    def test_real_insert_deduplicate_versions_and_query(self):
        preview = service.stage_uploads(list(self.files.items()), 'tester', self.root)
        result = service.persist_upload(preview['token'], 'tester', root=self.root, session_factory=self.factory)
        self.assertEqual(result['inserted'], 96)
        first = self.connection.execute(text(f"SELECT * FROM {service.TABLE} WHERE trade_date='2026-01-12' AND point_index=1")).mappings().one()
        self.assertEqual(first['bill_real_time_price'], -100)
        self.assertEqual(first['xianghai_unit2_sales'], 0)
        for report in [service.parse_bill(content, name) for name, content in self.files.items()]:
            for key, expected in service.map_bill(report).items():
                actual = self.connection.execute(text(f'SELECT * FROM {service.TABLE} WHERE trade_date=:day AND point_index=:point'),
                                                 {'day': key[0], 'point': key[1]}).mappings().one()
                for field, value in expected.items():
                    self.assertEqual(actual[field], value, field)
        again = service.persist_upload(preview['token'], 'tester', root=self.root, session_factory=self.factory)
        self.assertEqual((again['inserted'], len(again['skipped'])), (0, 2))
        # 单厂更正需确认；另一厂字段及公开实时价不变。
        changed = edited(self.files[NAMES[0]], lambda b: setattr(b.active['F11'], 'value', None))
        revision = service.stage_uploads([(NAMES[0], changed)], 'tester', self.root)
        revised = service.persist_upload(revision['token'], 'tester', root=self.root, session_factory=self.factory)
        self.assertTrue(revised['needs_overwrite'])
        revised = service.persist_upload(revision['token'], 'tester', root=self.root, session_factory=self.factory, overwrite=True)
        self.assertEqual(revised['inserted'], 96)
        current = self.connection.execute(text(f"SELECT * FROM {service.TABLE} WHERE trade_date='2026-01-12' AND point_index=1")).mappings().one()
        self.assertIsNone(current['xianghai_unit1_contract_quantity'])
        self.assertEqual(current['beihai_unit1_sales'], first['beihai_unit1_sales'])
        self.assertEqual(current['real_time_price'], first['real_time_price'])
        app = FastAPI()
        app.include_router(api.router)
        app.dependency_overrides[api._importer] = lambda: SimpleNamespace(username='tester', group='Global_admin')
        with patch.object(api, 'SessionLocal', self.factory), patch.object(service, 'IMPORT_ROOT', self.root):
            client = TestClient(app)
            history = client.get('/spot_electricity/bills/imports').json()
            self.assertEqual(len(history['items']), 2)
            import_id = revised['imported'][0]['import_id']
            details = client.get('/spot_electricity/bills/records', params={'import_id': import_id, 'limit': 3}).json()
            self.assertEqual(len(details['items']), 3)
            self.assertEqual(details['items'][0]['time_label'], '00:15')
            self.assertEqual(details['total'], 4 * 96 * 2)
            self.assertIn('cell', details['items'][0]['source_cells']['quantity'])

    def test_shared_price_conflict_does_not_modify_data(self):
        preview = service.stage_uploads([(NAMES[0], self.files[NAMES[0]])], 'tester', self.root)
        service.persist_upload(preview['token'], 'tester', root=self.root, session_factory=self.factory)
        # 修改北海所有机组同一时段价格，厂内一致但与已导入香海不一致。
        def change(book):
            for row in book.active:
                if row[0].value == '00:15':
                    book.active.cell(row[0].row + 1, 7).value = 999
        changed = edited(self.files[NAMES[1]], change)
        preview = service.stage_uploads([(NAMES[1], changed)], 'tester', self.root)
        with self.assertRaisesRegex(ValueError, '另一电厂'):
            service.persist_upload(preview['token'], 'tester', root=self.root, session_factory=self.factory, overwrite=True)
        row = self.connection.execute(text(f"SELECT bill_real_time_price,beihai_bill_import_id FROM {service.TABLE} WHERE trade_date='2026-01-12' AND point_index=1")).one()
        self.assertEqual(row.bill_real_time_price, -100)
        self.assertIsNone(row.beihai_bill_import_id)

    def test_database_failure_rolls_back_entire_batch(self):
        preview = service.stage_uploads(list(self.files.items()), 'tester', self.root)
        before = self.connection.execute(text(f'SELECT count(*) FROM {service.TABLE}')).scalar()
        before_values = self.connection.execute(text(f"SELECT * FROM {service.TABLE} WHERE trade_date='2026-01-12' ORDER BY point_index")).mappings().all()
        inserts = []
        def fail_second_insert(conn, cursor, statement, parameters, context, executemany):
            if statement.startswith('INSERT INTO ' + service.TABLE):
                inserts.append(statement)
                if len(inserts) == 2:
                    raise RuntimeError('模拟中途写入失败')
        event.listen(self.connection, 'before_cursor_execute', fail_second_insert)
        try:
            with self.assertRaisesRegex(RuntimeError, '中途写入失败'):
                service.persist_upload(preview['token'], 'tester', root=self.root, session_factory=self.factory)
        finally:
            event.remove(self.connection, 'before_cursor_execute', fail_second_insert)
        after = self.connection.execute(text(f'SELECT count(*) FROM {service.TABLE}')).scalar()
        self.assertEqual(after, before)
        after_values = self.connection.execute(text(f"SELECT * FROM {service.TABLE} WHERE trade_date='2026-01-12' ORDER BY point_index")).mappings().all()
        self.assertEqual(after_values, before_values)


if __name__ == '__main__':
    if '--stdin-fixtures' in sys.argv:
        sys.argv.remove('--stdin-fixtures')
        FIXTURES = {name: base64.b64decode(value) for name, value in json.load(sys.stdin).items()}
    if '--database' in sys.argv:
        sys.argv.remove('--database')
        os.environ['SPOT_BILL_TEST_DB'] = '1'
    unittest.main()
