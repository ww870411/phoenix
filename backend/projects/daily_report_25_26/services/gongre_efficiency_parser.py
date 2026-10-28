"""
供热公司能效分析表（Excel）解析服务
严格仅抽取【本期】生产数据，严禁触碰或改变同期值数据。
对应生产日报子表：
1. GongRe_Sheet: 5.供热公司表
2. GongRe_branches_detail_Sheet: 6.供热公司分中心明细表（10个分中心换热站补水量）
"""

import io
import datetime
import re
from typing import Dict, Any, Optional
import openpyxl
from openpyxl.utils.datetime import from_excel


def _clean_val(v: Any) -> Optional[float]:
    """清洗数值：支持过滤异常公式错误字符与空值，规范为浮点数"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 4)
    s = str(v).strip()
    if not s or s.startswith('#') or s in ('-', 'None', 'null'):
        return None
    try:
        s_clean = s.replace(',', '')
        return round(float(s_clean), 4)
    except (ValueError, TypeError):
        return None


def _clean_str(v: Any) -> str:
    """清洗文本"""
    if v is None:
        return ""
    s = str(v).strip()
    if s.startswith('#') or s in ('None', 'null'):
        return ""
    return s


def parse_gongre_efficiency_excel(
    file_bytes: bytes,
    biz_date_str: str,
) -> Dict[str, Any]:
    """
    解析供热公司能效分析表 Excel 文件，严格仅提取【本期】（当前业务年份）生产数据。
    严禁提取或覆盖任何历史同期数据。
    """
    try:
        dt = datetime.datetime.strptime(biz_date_str.strip(), "%Y-%m-%d")
    except ValueError:
        raise ValueError(f"无效的业务日期格式: {biz_date_str}，应为 YYYY-MM-DD")

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    target_year = dt.year
    target_month = dt.month
    target_day = dt.day

    # 1. 寻找匹配的工作表 (Sheet)
    sheet_name = None
    for s in wb.sheetnames:
        clean_name = s.strip()
        if clean_name in (
            f"{target_year}.{target_month}",
            f"{target_year}.{target_month:02d}",
            f"{target_year}年{target_month}月",
            f"供热{target_year}.{target_month}",
            f"供热{target_year}.{target_month:02d}",
        ):
            sheet_name = s
            break

    if not sheet_name:
        for s in wb.sheetnames:
            clean_name = s.strip()
            if str(target_year) in clean_name and f".{target_month}" in clean_name:
                sheet_name = s
                break

    if not sheet_name:
        raise ValueError(
            f"未在上传的 Excel 中找到业务月份对应的工作表（需匹配 {target_year}.{target_month}），当前包含的工作表为: {', '.join(wb.sheetnames)}"
        )

    ws = wb[sheet_name]

    # 2. 在 Row 2 中寻找日期对应的列 (从第 3 列开始)
    date_col = None
    for col in range(3, ws.max_column + 1):
        cell_val = ws.cell(2, col).value
        if cell_val is not None:
            try:
                if isinstance(cell_val, (datetime.datetime, datetime.date)):
                    if cell_val.day == target_day:
                        date_col = col
                        break
                elif isinstance(cell_val, (int, float)):
                    dt_val = from_excel(cell_val)
                    if dt_val.day == target_day:
                        date_col = col
                        break
                else:
                    m = re.match(r"^(\d+)", str(cell_val).strip())
                    if m and int(m.group(1)) == target_day:
                        date_col = col
                        break
            except Exception:
                pass

    if not date_col:
        for r in (1, 3):
            for col in range(3, ws.max_column + 1):
                cell_val = ws.cell(r, col).value
                if cell_val is not None:
                    try:
                        if isinstance(cell_val, (datetime.datetime, datetime.date)):
                            if cell_val.day == target_day:
                                date_col = col
                                break
                        elif isinstance(cell_val, (int, float)):
                            dt_val = from_excel(cell_val)
                            if dt_val.day == target_day:
                                date_col = col
                                break
                        else:
                            m = re.match(r"^(\d+)", str(cell_val).strip())
                            if m and int(m.group(1)) == target_day:
                                date_col = col
                                break
                    except Exception:
                        pass
            if date_col:
                break

    if not date_col:
        raise ValueError(
            f"在工作表【{sheet_name}】中未找到业务日期 {target_day} 日对应的列"
        )

    # 3. 抽取本期数据
    def get_num(row_idx: int) -> float:
        """获取指定行数值，空值或不可解析返回 0.0"""
        v = ws.cell(row_idx, date_col).value
        cleaned = _clean_val(v)
        return cleaned if cleaned is not None else 0.0

    # === 表 5: 供热公司表 (GongRe_Sheet) ===
    # 外购电量：Row 29 站内耗电总量(kwh)，日报单位为万kWh，换算除以 10000.0
    station_power_kwh = get_num(29)
    station_power_wan_kwh = round(station_power_kwh / 10000.0, 4) if station_power_kwh else 0.0

    # 补水量：
    # Row 24: 耗水总量 -> 日报系统【耗水量】
    # Row 23: 2026年电厂侧 -> 日报系统【其中：电厂一次网补水量】
    # Row 22: 2026年站内 -> 日报系统【其中：换热站补水量】
    water_total = get_num(24)
    water_plant_side = get_num(23)
    water_station = get_num(22)

    # 气温：
    # Row 6: 2026 当日最低温度
    # Row 7: 2026 当日最高温度
    temp_min = get_num(6)
    temp_max = get_num(7)

    gongre_cells = [
        {"row_label": "网口供热量", "unit": "GJ", "value": get_num(10)},
        {"row_label": "高温水销售量", "unit": "GJ", "value": 0.0},
        {"row_label": "售汽量", "unit": "吨", "value": get_num(40)},
        {"row_label": "耗水量", "unit": "吨", "value": water_total},
        {"row_label": "其中：电厂一次网补水量", "unit": "吨", "value": water_plant_side},
        {"row_label": "其中：换热站补水量", "unit": "吨", "value": water_station},
        {"row_label": "外购电量", "unit": "万kWh", "value": station_power_wan_kwh},
        {"row_label": "其中：换热站外购电量", "unit": "万kWh", "value": station_power_wan_kwh},
        {"row_label": "站内耗热量", "unit": "GJ", "value": get_num(15)},
        # 售高温水收入与售汽收入遵循业务要求不做自动识别，留空
        {"row_label": "日最高气温", "unit": "℃", "value": temp_max},
        {"row_label": "日最低气温", "unit": "℃", "value": temp_min},
    ]

    # === 表 6: 供热公司分中心明细表 (GongRe_branches_detail_Sheet) ===
    # 动态扫描能效报表下部列出各中心简称、指标名为“耗水量”的行（如“华乐  耗水量（吨）”）
    # 在分中心明细表中对应指标“其中：换热站补水量”
    # 分中心明细表列定义：c0=项目, c1=中心, c2=计量单位, c3=(本期日), c4=(同期日), c5=解释说明
    center_definitions = [
        ("华乐中心", "华乐", 188),
        ("东港中心", "东港", 189),
        ("春海中心", "春海", 190),
        ("中山中心", "中山", 191),
        ("北岗中心", "北岗", 192),
        ("双兴中心", "双兴", 193),
        ("香炉礁中心", "香炉礁", 194),
        ("兴工中心", "兴工", 195),
        ("大纺中心", "大纺", 196),
        ("管网中心", "管网", 197),
    ]

    # 动态嗅探定位行号，优先使用动态匹配到的行号，以默认行号兜底
    dynamic_branch_rows = {}
    for r in range(150, ws.max_row + 1):
        a_val = str(ws.cell(r, 1).value or '')
        if '耗水' in a_val:
            for full_name, short_name, _ in center_definitions:
                if short_name in a_val and full_name not in dynamic_branch_rows:
                    dynamic_branch_rows[full_name] = r
                    break

    branches_cells = []
    for full_name, short_name, default_row in center_definitions:
        row_idx = dynamic_branch_rows.get(full_name, default_row)
        val = get_num(row_idx)
        branches_cells.append({
            "row_label": "其中：换热站补水量",
            "sub_label": full_name,
            "unit": "吨",
            "col_index": 3,  # 本期日所在列索引为 3
            "value": val,
        })

    total_matched = len(gongre_cells) + len(branches_cells)

    return {
        "ok": True,
        "biz_date": biz_date_str,
        "sheet_name": sheet_name,
        "day": target_day,
        "data": {
            "GongRe_Sheet": {
                "sheet_name": "5.供热公司表",
                "matched_count": len(gongre_cells),
                "cells": gongre_cells,
            },
            "GongRe_branches_detail_Sheet": {
                "sheet_name": "6.供热公司分中心明细表",
                "matched_count": len(branches_cells),
                "cells": branches_cells,
            },
        },
        "summary": {
            "total_matched": total_matched,
            "message": f"成功从工作表【{sheet_name}】第 {target_day} 日列提取供热公司总表（{len(gongre_cells)} 项）及分中心补水量明细（{len(branches_cells)} 项）【本期实际值】（同期值保持锁定未修改）！"
        }
    }
