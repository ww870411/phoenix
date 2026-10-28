"""
香海热电厂能效分析表（Excel）解析服务
严格仅抽取【本期】生产数据，严禁触碰或改变同期值数据。
对应生产日报子表：
XiangHai_Sheet: 4.香海热电厂表
"""

import io
import datetime
import re
from typing import Dict, Any, Optional, Tuple
import openpyxl
from openpyxl.utils.datetime import from_excel


def _clean_val(v: Any) -> Optional[float]:
    """清洗数值：支持过滤异常公式错误字符与空值，规范为浮点数"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 4)
    s = str(v).strip()
    if not s or s.startswith('#') or s in ('-', 'None', 'null'):
        return None
    try:
        s_clean = s.replace(',', '')
        return round(float(s_clean), 4)
    except (ValueError, TypeError):
        return None


def _clean_str(v: Any) -> str:
    """清洗文本，过滤 Excel 常见错误码"""
    if v is None:
        return ""
    s = str(v).strip()
    # 仅过滤 Excel 错误公式字符串（如 #VALUE!, #DIV/0!, #REF!），不得过滤 #1炉、#2机 等设备工况
    if re.match(r"^#[A-Z/!?]+", s) or s in ('None', 'null'):
        return ""
    return s


def _parse_steam_boiler_and_turbine(s: str) -> Tuple[int, int]:
    """解析香海运行方式文本，获取运行汽炉数与运行汽轮机数。若为空、停运或无记录则识别为 (0, 0)"""
    if not s or s in ('0', '停', '无', 'None'):
        return 0, 0
    s_clean = s.replace('\n', ' ').strip()
    if not s_clean or s_clean in ('0', '停', '无'):
        return 0, 0

    # 查找汽炉数
    boiler_cnt = 0
    m_boiler = re.search(r'([^机]*?)炉', s_clean)
    if m_boiler:
        nums = re.findall(r'\d+', m_boiler.group(1))
        boiler_cnt = len(nums) if nums else 1
    elif '炉' in s_clean:
        boiler_cnt = 1

    # 查找汽轮机数
    turbine_cnt = 0
    m_turb = re.search(r'([^炉]*?)机', s_clean)
    if m_turb:
        nums = re.findall(r'\d+', m_turb.group(1))
        turbine_cnt = len(nums) if nums else 1
    elif '机' in s_clean:
        turbine_cnt = 1

    return boiler_cnt, turbine_cnt


def parse_xianghai_efficiency_excel(
    file_bytes: bytes,
    biz_date_str: str,
) -> Dict[str, Any]:
    """
    解析香海能效分析表 Excel 文件，严格仅提取【本期】（当前业务年份）生产数据。
    严禁提取或覆盖任何历史同期数据。
    """
    try:
        dt = datetime.datetime.strptime(biz_date_str.strip(), "%Y-%m-%d")
    except ValueError:
        raise ValueError(f"无效的业务日期格式: {biz_date_str}，应为 YYYY-MM-DD")

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    target_year = dt.year
    target_month = dt.month
    target_day = dt.day

    # 1. 寻找匹配的工作表 (Sheet)
    sheet_name = None
    # 候选模式（按优先级匹配，去除空格对比）
    for s in wb.sheetnames:
        clean_name = s.strip()
        # 匹配香海YYYY.M 或 YYYY.M
        if clean_name in (
            f"香海{target_year}.{target_month}",
            f"香海{target_year}.{target_month:02d}",
            f"{target_year}.{target_month}",
            f"{target_year}.{target_month:02d}",
        ):
            sheet_name = s
            break

    if not sheet_name:
        for s in wb.sheetnames:
            clean_name = s.strip()
            if str(target_year) in clean_name and f".{target_month}" in clean_name:
                sheet_name = s
                break

    if not sheet_name:
        raise ValueError(
            f"未在上传的 Excel 中找到业务月份对应的工作表（需匹配 {target_year}.{target_month}），当前包含的工作表为: {', '.join(wb.sheetnames)}"
        )

    ws = wb[sheet_name]

    # 2. 在 Row 2 中寻找日期对应的列 (从第 3 列开始)
    date_col = None
    for col in range(3, ws.max_column + 1):
        cell_val = ws.cell(2, col).value
        if cell_val is not None:
            try:
                if isinstance(cell_val, (datetime.datetime, datetime.date)):
                    if cell_val.day == target_day:
                        date_col = col
                        break
                elif isinstance(cell_val, (int, float)):
                    # Excel 序列号
                    dt_val = from_excel(cell_val)
                    if dt_val.day == target_day:
                        date_col = col
                        break
                else:
                    m = re.match(r"^(\d+)", str(cell_val).strip())
                    if m and int(m.group(1)) == target_day:
                        date_col = col
                        break
            except Exception:
                pass

    if not date_col:
        # 兜底在 Row 1 或 Row 3 中查找
        for r in (1, 3):
            for col in range(3, ws.max_column + 1):
                cell_val = ws.cell(r, col).value
                if cell_val is not None:
                    try:
                        if isinstance(cell_val, (datetime.datetime, datetime.date)):
                            if cell_val.day == target_day:
                                date_col = col
                                break
                        elif isinstance(cell_val, (int, float)):
                            dt_val = from_excel(cell_val)
                            if dt_val.day == target_day:
                                date_col = col
                                break
                        else:
                            m = re.match(r"^(\d+)", str(cell_val).strip())
                            if m and int(m.group(1)) == target_day:
                                date_col = col
                                break
                    except Exception:
                        pass
            if date_col:
                break

    if not date_col:
        raise ValueError(
            f"在工作表【{sheet_name}】中未找到业务日期 {target_day} 日对应的列"
        )

    # 3. 抽取本期数据（只取本期当年专属行，严禁触碰同期行）
    def get_num(row_idx: int) -> float:
        """获取指定行数值，空值或异常返回 0.0"""
        v = ws.cell(row_idx, date_col).value
        cleaned = _clean_val(v)
        return cleaned if cleaned is not None else 0.0

    def get_str(row_idx: int) -> str:
        """获取指定行文本"""
        return _clean_str(ws.cell(row_idx, date_col).value)

    # 解析机炉运行方式（Row 10 为当年本期）
    run_status_str = get_str(10)
    boiler_count, turbine_count = _parse_steam_boiler_and_turbine(run_status_str)

    # 外购电量：能效表 Row 89 单位为千瓦时(kWh)，日报系统单位为万kWh，换算除以 10000.0
    power_kwh = get_num(89)
    power_wan_kwh = round(power_kwh / 10000.0, 4) if power_kwh else 0.0

    # 耗水量说明：
    # Row 73 为自来水量 -> 对应日报系统【耗水量】
    # Row 76 为电厂耗水量 -> 对应日报系统【其中：电厂耗水量】
    water_total = get_num(73)
    water_plant = get_num(76)

    cells = [
        {"row_label": "发电量", "unit": "万kWh", "value": get_num(13)},
        {"row_label": "供热量", "unit": "GJ", "value": get_num(50)},
        {"row_label": "供汽量", "unit": "吨", "value": get_num(43)},
        {"row_label": "售电量", "unit": "万kWh", "value": get_num(28)},
        {"row_label": "原煤耗量", "unit": "吨", "value": get_num(58)},
        {"row_label": "标煤耗量", "unit": "吨", "value": get_num(63)},
        {"row_label": "耗水量", "unit": "吨", "value": water_total},
        {"row_label": "其中：电厂耗水量", "unit": "吨", "value": water_plant},
        {"row_label": "外购电量", "unit": "万kWh", "value": power_wan_kwh},
        {"row_label": "其中：电厂外购电量", "unit": "万kWh", "value": power_wan_kwh},
        {"row_label": "耗酸量", "unit": "吨", "value": get_num(109)},
        {"row_label": "耗碱量", "unit": "吨", "value": get_num(114)},
        {"row_label": "耗油量", "unit": "吨", "value": get_num(84)},
        {"row_label": "耗氨水量", "unit": "吨", "value": get_num(94)},
        {"row_label": "耗石灰石粉量", "unit": "吨", "value": get_num(99)},
        {"row_label": "运行汽炉数", "unit": "台", "value": boiler_count},
        {"row_label": "运行汽轮机数", "unit": "台", "value": turbine_count},
    ]

    return {
        "ok": True,
        "biz_date": biz_date_str,
        "sheet_name": sheet_name,
        "day": target_day,
        "data": {
            "XiangHai_Sheet": {
                "sheet_name": "4.香海热电厂表",
                "matched_count": len(cells),
                "cells": cells,
            }
        },
        "summary": {
            "total_matched": len(cells),
            "message": f"成功从工作表【{sheet_name}】第 {target_day} 日列提取香海热电厂【本期实际值】共 {len(cells)} 项（同期值保持锁定未修改）！"
        }
    }
