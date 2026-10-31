"""原始电力日账单解析和事务入库；不推算、不补零、不筛除未运行机组。"""
from __future__ import annotations

import hashlib
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import UUID, uuid4

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import text

from backend.config import DATA_DIRECTORY
from backend.db.database_daily_report_25_26 import SessionLocal

IMPORT_ROOT = DATA_DIRECTORY / 'projects' / 'daily_report_25_26' / 'imports' / 'spot_settlement'
TABLE = 'public.spot_electricity_15min_data'
PLANTS = {'BeiHai': ('beihai', 3), 'XiangHai': ('xianghai', 4)}


def _text(value):
    return '' if value is None else str(value).strip()


def _json_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else str(value) if isinstance(value, Decimal) else value


def parse_bill(content: bytes, filename: str) -> dict:
    """解析所有非空工作表；缺时段、重复时段或缺少身份信息时拒绝整份文件。"""
    digest = hashlib.sha256(content).hexdigest()
    name = Path(filename.replace('\\', '/')).name
    formula_book = load_workbook(io.BytesIO(content), data_only=False)
    value_book = load_workbook(io.BytesIO(content), data_only=True)
    records, units, warnings = [], [], []
    try:
        for sheet in formula_book:
            values = value_book[sheet.title]
            if not any(cell.value is not None for row in sheet for cell in row):
                continue

            def cell_info(row, col):
                c, v = sheet.cell(row, col), values.cell(row, col)
                return {'cell': c.coordinate, 'value': _json_value(c.value),
                        'cached_value': _json_value(v.value), 'number_format': c.number_format, 'data_type': c.data_type}

            def numeric(row, col):
                value = values.cell(row, col).value
                if value is None or value == '':
                    if sheet.cell(row, col).data_type == 'f':
                        warnings.append(f'{sheet.title}!{sheet.cell(row, col).coordinate} 公式无缓存，已保留公式')
                    return None
                try:
                    result = Decimal(str(value).strip())
                    if not result.is_finite():
                        raise InvalidOperation
                    return result
                except (InvalidOperation, ValueError):
                    warnings.append(f'{sheet.title}!{sheet.cell(row, col).coordinate} 非数值原文已保留')
                    return None

            def find_label(label, rows):
                for row in rows:
                    for col in range(1, sheet.max_column + 1):
                        if _text(sheet.cell(row, col).value).rstrip('：: ') == label:
                            for nxt in range(col + 1, sheet.max_column + 1):
                                if sheet.cell(row, nxt).value is not None:
                                    return row, nxt
                raise ValueError(f'{sheet.title} 缺少{label}')

            def header(label):
                row, col = find_label(label, range(1, min(sheet.max_row, 4) + 1))
                return values.cell(row, col).value

            settlement_date = header('结算日期')
            if isinstance(settlement_date, (int, float)):
                settlement_date = from_excel(settlement_date, formula_book.epoch)
            if isinstance(settlement_date, datetime):
                settlement_date = settlement_date.date()
            elif not isinstance(settlement_date, date):
                settlement_date = date.fromisoformat(_text(settlement_date)[:10])
            title_date = re.search(r'\d{4}-\d{2}-\d{2}', _text(sheet.cell(1, 1).value))
            if title_date and title_date[0] != settlement_date.isoformat():
                raise ValueError(f'{sheet.title} 标题日期与结算日期不一致')
            member, bill = _text(header('市场成员名称')), _text(header('账单编号'))
            if not member or not bill:
                raise ValueError('市场成员名称及账单编号不能为空')
            qty_pos, fee_pos = find_label('合计电量', range(1, 5)), find_label('合计电费', range(1, 5))
            common = dict(file_sha256=digest, source_file=name, worksheet=sheet.title,
                          settlement_date=settlement_date, bill_number=bill, member_name=member,
                          generator_account=_text(header('发电户编号')),
                          unit_label=' '.join(_text(cell.value) for row in sheet.iter_rows(min_row=1, max_row=4)
                                              for cell in row if _text(cell.value).startswith(('单位:', '单位：'))),
                          bill_total_quantity=numeric(*qty_pos), bill_total_fee=numeric(*fee_pos))
            starts = [row for row in range(1, sheet.max_row + 1)
                      if _text(sheet.cell(row, 1).value).startswith('结算单元名称')]
            if not starts:
                raise ValueError(f'{sheet.title} 未识别到结算单元，未导入')
            identities = set()
            for index, start in enumerate(starts):
                end = starts[index + 1] if index + 1 < len(starts) else sheet.max_row + 1
                unit_name = re.sub(r'^结算单元名称\s*[:：]\s*', '', _text(sheet.cell(start, 1).value))
                id_cells = [_text(sheet.cell(start, col).value) for col in range(1, sheet.max_column + 1)]
                ids = [re.sub(r'^结算单元编号\s*[:：]\s*', '', v) for v in id_cells if v.startswith('结算单元编号')]
                if not ids or not ids[0] or ids[0] in identities:
                    raise ValueError(f'{sheet.title} 结算单元编号缺失或重复')
                identities.add(ids[0])
                number = re.search(r'(\d+)号机', unit_name)
                plant_name = re.sub(r'\d+号机.*$', '', unit_name).strip()
                plant_key = 'BeiHai' if '北海' in plant_name else 'XiangHai' if '香海' in plant_name else plant_name
                unit_qty = find_label('结算单元计量电量', [start])
                unit_fee = find_label('结算单元小计电费', [start])
                base = dict(common, plant_key=plant_key, plant_name=plant_name,
                            settlement_unit_id=ids[0], settlement_unit_name=unit_name,
                            unit_number=int(number[1]) if number else None,
                            unit_metered_quantity=numeric(*unit_qty), unit_total_fee=numeric(*unit_fee))
                code_rows = [r for r in range(start + 1, end) if _text(sheet.cell(r, 2).value) == '科目编码']
                if len(code_rows) != 1:
                    raise ValueError(f'{unit_name} 科目表头缺失或重复')
                code_row = code_rows[0]
                subjects = []
                for col in range(3, sheet.max_column + 1):
                    raw_code = sheet.cell(code_row, col).value
                    # Excel 数值编码可能读成100600.0，身份键规范为100600；原单元格另存。
                    code = str(int(raw_code)) if isinstance(raw_code, (int, float)) and not isinstance(raw_code, bool) and raw_code == int(raw_code) else _text(raw_code)
                    if not code:
                        continue
                    subject_name = _text(sheet.cell(code_row + 1, col).value)
                    if not subject_name or code in [item[1] for item in subjects]:
                        raise ValueError(f'{unit_name} 科目名称缺失或编码重复')
                    subjects.append((col, code, subject_name))
                if not subjects:
                    raise ValueError(f'{unit_name} 未识别到科目')
                point_rows, current_point = {}, None
                for row in range(code_row + 2, end):
                    label = _text(sheet.cell(row, 1).value)
                    if label == '合计':
                        current_point = 0
                    elif re.fullmatch(r'\d{2}:\d{2}', label):
                        hour, minute = map(int, label.split(':'))
                        total = hour * 60 + minute
                        if minute not in (0, 15, 30, 45) or not 15 <= total <= 1440 or hour > 24 or (hour == 24 and minute):
                            raise ValueError(f'{unit_name} 非法时段 {label}')
                        current_point = total // 15
                        if current_point in point_rows:
                            raise ValueError(f'{unit_name} 重复时段 {label}')
                    elif label:
                        raise ValueError(f'{unit_name} 无法识别的时段标签：{label}')
                    metric = _text(sheet.cell(row, 2).value)
                    if metric not in ('电量', '电价', '电费'):
                        if any(sheet.cell(row, c).value is not None for c in range(1, sheet.max_column + 1)):
                            raise ValueError(f'{unit_name} 第{row}行内容无法识别，未跳过入库')
                        continue
                    if current_point is None:
                        raise ValueError(f'{unit_name} 第{row}行无所属时段')
                    bucket = point_rows.setdefault(current_point, {})
                    if metric in bucket:
                        raise ValueError(f'{unit_name} 同时段重复{metric}')
                    bucket[metric] = row
                if set(point_rows) != set(range(97)):
                    raise ValueError(f'{unit_name} 必须包含合计及完整96个时段')
                for point, metric_rows in point_rows.items():
                    needed = {'电量', '电费'} if point == 0 else {'电量', '电价', '电费'}
                    if not needed.issubset(metric_rows):
                        raise ValueError(f'{unit_name} 第{point}时段量价费行不完整')
                    for col, code, subject in subjects:
                        positions = {field: (metric_rows[kind], col) for field, kind in
                                     [('quantity', '电量'), ('price', '电价'), ('fee', '电费')] if kind in metric_rows}
                        records.append(dict(base, record_type='period' if point else 'daily_total',
                                            point_index=point, time_label=f'{point // 4:02d}:{point % 4 * 15:02d}' if point else '合计',
                                            subject_code=code, subject_name=subject,
                                            **{field: numeric(*positions[field]) if field in positions else None for field in ('quantity', 'price', 'fee')},
                                            source_cells={field: cell_info(*position) for field, position in positions.items()} |
                                            {'subject_code': cell_info(code_row, col), 'subject_name': cell_info(code_row + 1, col)}))
                units.append(dict(base, points=96, subjects=len(subjects), records=97 * len(subjects)))
        if not records:
            raise ValueError('工作簿没有可识别的日账单')
        return {'file_sha256': digest, 'source_file': name, 'records': records, 'units': units,
                'warnings': list(dict.fromkeys(warnings))}
    finally:
        formula_book.close()
        value_book.close()


def stage_uploads(files, owner, root=None):
    """预览只归档原文件，不写业务表；确认时从原文件重新解析。"""
    root = Path(root or IMPORT_ROOT)
    parsed = [parse_bill(content, name) for name, content in files]
    for report in parsed:
        map_bill(report)
    token = str(uuid4())
    folder = root / 'pending' / token
    folder.mkdir(parents=True, exist_ok=False)
    entries = []
    seen = set()
    for (name, content), result in zip(files, parsed):
        digest = result['file_sha256']
        if digest in seen:
            continue
        seen.add(digest)
        (folder / f'{digest}.xlsx').write_bytes(content)
        entries.append({'sha256': digest, 'filename': result['source_file']})
    (folder / 'manifest.json').write_text(json.dumps({'owner': owner, 'files': entries}, ensure_ascii=False), encoding='utf-8')
    unique_parsed = {}
    for report in parsed:
        unique_parsed.setdefault(report['file_sha256'], report)
    return {'ok': True, 'token': token, 'files': [
        {key: value for key, value in p.items() if key != 'records'} | {'record_count': len(p['records'])}
        for p in unique_parsed.values()]}


def map_bill(report):
    """按目标宽表映射计量电量、合约量价及共用实时价，不补零、不计算电费。"""
    identities = set()
    for unit in report['units']:
        plant = PLANTS.get(unit['plant_key'])
        identity = (unit['settlement_date'], unit['plant_key'], unit['unit_number'])
        if not plant or unit['unit_number'] not in range(1, plant[1] + 1) or identity in identities:
            raise ValueError('电厂或机组无法映射到目标表，或文件内机组重复')
        identities.add(identity)
    # 批次标识按电厂保存，因此每份文件必须覆盖该厂全部机组。
    for day, plant, _ in identities:
        if {n for d, p, n in identities if (d, p) == (day, plant)} != set(range(1, PLANTS[plant][1] + 1)):
            raise ValueError('日账单必须包含该电厂全部机组，包括未运行机组')
    rows, prices, seen = {}, {}, set()
    for record in report['records']:
        if record['record_type'] != 'period' or record['subject_code'] not in ('100401', '100600'):
            continue
        day, point = record['settlement_date'], record['point_index']
        key = (day, point)
        prefix = f"{PLANTS[record['plant_key']][0]}_unit{record['unit_number']}"
        seen.add((day, record['plant_key'], record['unit_number'], point, record['subject_code']))
        row = rows.setdefault(key, {'trade_date': day, 'point_index': point, 'time_label': record['time_label'],
                                  'period_start_time': f'{(point - 1) // 4:02d}:{(point - 1) % 4 * 15:02d}:00',
                                  'period_end_time': record['time_label'] + ':00'})
        if record['subject_code'] == '100401':
            row[prefix + '_contract_quantity'] = record['quantity']
            row[prefix + '_contract_price'] = record['price']
        else:
            quantity = record['quantity']
            if quantity is not None and (abs(quantity) >= Decimal('100000000') or quantity != quantity.quantize(Decimal('0.0001'))):
                raise ValueError(f'{day} {record["settlement_unit_name"]} {record["time_label"]} 计量电量超出目标字段NUMERIC(12,4)精度，未四舍五入入库')
            row[prefix + '_sales'] = record['quantity']
            if key in prices and prices[key] != record['price']:
                raise ValueError(f'{day} {record["time_label"]} 各机组省内实时电价不一致，不能合并为共用价格')
            prices[key] = record['price']
            row['bill_real_time_price'] = record['price']
    expected = {(d, p, n, point, code) for d, p, n in identities for point in range(1, 97) for code in ('100401', '100600')}
    if seen != expected:
        raise ValueError('缺少完整96时段的省内实时或省内中长期合约科目')
    return rows


def load_batch(import_id, root=None):
    """原文件及来源清单保存在持久化数据目录；数据库只保留当前电厂批次标识。"""
    path = Path(root or IMPORT_ROOT) / 'batches' / f'{UUID(str(import_id))}.json'
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else {}


def persist_upload(token, owner, allowed_units=None, root=None, session_factory=None, overwrite=False):
    """日期锁保护共用价格；整批事务更新，仅覆盖本次电厂字段。"""
    root = Path(root or IMPORT_ROOT)
    folder = root / 'pending' / str(UUID(str(token)))
    manifest = json.loads((folder / 'manifest.json').read_text(encoding='utf-8'))
    if manifest['owner'] != owner:
        raise PermissionError('只能确认本人上传的预览')
    parsed, identities = [], set()
    for entry in manifest['files']:
        content = (folder / f"{entry['sha256']}.xlsx").read_bytes()
        report = parse_bill(content, entry['filename'])
        if report['file_sha256'] != entry['sha256']:
            raise ValueError('预览文件内容发生变化，请重新上传')
        if allowed_units is not None and any(u['plant_key'] not in allowed_units for u in report['units']):
            raise PermissionError('上传文件包含无权提交的电厂')
        mapped = map_bill(report)
        current = {(u['settlement_date'], u['plant_key']) for u in report['units']}
        if identities & current:
            raise ValueError('本批包含同电厂同日的多份账单，请只保留一个版本')
        identities.update(current)
        parsed.append((report, content, mapped))
    prices = {}
    for report, content, mapped in parsed:
        for key, row in mapped.items():
            if key in prices and prices[key] != row['bill_real_time_price']:
                raise ValueError(f'{key[0]} {row["time_label"]} 两厂账单实时价不一致')
            prices[key] = row['bill_real_time_price']
    imported, skipped, conflicts, combined = [], [], [], {}
    with (session_factory or SessionLocal)() as db:
        with db.begin():
            existing = {}
            for day in sorted({d for d, _ in identities}):
                db.execute(text('SELECT pg_advisory_xact_lock(:key)'), {'key': 1500000000 + day.toordinal()})
                for row in db.execute(text(f'SELECT * FROM {TABLE} WHERE trade_date=:day FOR UPDATE'), {'day': day}).mappings():
                    existing[(day, row['point_index'])] = row
            for report, content, mapped in parsed:
                batch_columns = {day: [PLANTS[p][0] + '_bill_import_id' for p in
                                 {u['plant_key'] for u in report['units'] if u['settlement_date'] == day}]
                                 for day, _ in mapped}
                same = all(key in existing and all(existing[key][field] == value for field, value in row.items()
                           if field.endswith(('_sales', '_quantity', '_price'))) and
                           all(load_batch(existing[key][c], root).get('file_sha256') == report['file_sha256']
                               for c in batch_columns[key[0]] if existing[key][c]) and all(existing[key][c] for c in batch_columns[key[0]])
                           for key, row in mapped.items())
                if same:
                    skipped.append(report['source_file'])
                    continue
                if any(key in existing and (any(existing[key][c] for c in batch_columns[key[0]]) or
                       any(existing[key][f] is not None for f in row if f.endswith(('_sales', '_quantity'))))
                       for key, row in mapped.items()):
                    conflicts.append(f'{report["source_file"]}：同日该电厂已有数据，将替换该电厂全部机组量价')
                import_id = str(uuid4())
                for key, row in mapped.items():
                    target = combined.setdefault(key, {})
                    if 'bill_real_time_price' in target and target['bill_real_time_price'] != row['bill_real_time_price']:
                        raise ValueError(f'{key[0]} {row["time_label"]} 两厂账单实时价不一致')
                    target.update(row)
                    target.update({c: import_id for c in batch_columns[key[0]]})
                imported.append({'filename': report['source_file'], 'file_sha256': report['file_sha256'], 'import_id': import_id,
                                 'records': len(mapped), 'units': len(report['units'])})
            for key, row in combined.items():
                prior = existing.get(key)
                if prior:
                    for plant, (prefix, count) in PLANTS.items():
                        if (key[0], plant) not in identities and (prior[prefix + '_bill_import_id'] or
                            any(prior[f'{prefix}_unit{n}_sales'] is not None for n in range(1, count + 1))):
                            if prior['bill_real_time_price'] != row['bill_real_time_price']:
                                raise ValueError(f'{key[0]} {row["time_label"]} 与另一电厂已入库账单实时价冲突，请核对或同时上传两厂一致版本')
            if conflicts and not overwrite:
                return {'ok': False, 'needs_overwrite': True, 'conflicts': conflicts}
            # 文件先落盘，提交失败时不会产生可查询的导入记录；不删除历史原件。
            for report, content, mapped in parsed:
                batch = next((i for i in imported if i['file_sha256'] == report['file_sha256']), None)
                if not batch:
                    continue
                (root / 'originals').mkdir(parents=True, exist_ok=True)
                (root / 'batches').mkdir(parents=True, exist_ok=True)
                (root / 'originals' / f'{report["file_sha256"]}.xlsx').write_bytes(content)
                metadata = {k: v for k, v in report.items() if k != 'records'}
                metadata.update(imported_by=owner, imported_at=datetime.now().isoformat(), import_id=batch['import_id'])
                (root / 'batches' / f'{batch["import_id"]}.json').write_text(
                    json.dumps(metadata, default=_json_value, ensure_ascii=False), encoding='utf-8')
            for key, row in sorted(combined.items()):
                columns = list(row)
                updates = [f'{c}=EXCLUDED.{c}' for c in columns if c not in ('trade_date', 'point_index')]
                db.execute(text(f'INSERT INTO {TABLE} ({",".join(columns)}) VALUES ({",".join(":" + c for c in columns)}) '
                                f'ON CONFLICT (trade_date,point_index) DO UPDATE SET {",".join(updates)},updated_at=now()'), row)
    return {'ok': True, 'imported': imported, 'skipped': skipped, 'warnings': [], 'inserted': len(combined)}
