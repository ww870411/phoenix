# -*- coding: utf-8 -*-
r"""
主城区调度生产日报 (DOCX / PDF) 统一结构化提取引擎 (100+ 工业宽表全物理字段展开版)

【核心设计目标】：
1. 彻底废除固定表格索引，表头与特征语义自适应嗅探（自适应 DOCX 6 表与 PDF 8 表）；
2. 提取并展开全量 100+ 工业物理字段（气象、6首站36项、4蒸汽网口12项、燃煤进耗存25项、产量计划与实际20项、网口热量6项、初期补水6项）；
3. 结构化存根兼顾保留（机炉矩阵拓扑与微观抢修工作票八元组 JSONB）；
4. 同时输出扁平大宽表键名与分组树形结构，100% 满足数据库大宽表存储与前端页面交互。
"""

import os
import re
import json
import logging
from .dispatch_reconciliation import reconcile_report
from typing import Dict, Any, List, Optional, Tuple
from datetime import datetime

logger = logging.getLogger(__name__)

try:
    import docx
except ImportError:
    docx = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None


class UnifiedDispatchExtractor:
    """通用调度生产日报解析器（自适应表头语义嗅探与企业拓扑标准化）"""

    KNOWN_ORGS = [
        ("北海电厂", ["北海电厂", "北海", "北海厂", "北海热电", "北海热电联产", "北海水炉", "燃气锅炉"]),
        ("香海电厂", ["香海电厂", "香海", "香海厂", "香海热电"]),
        ("金州热电", ["金州热电", "金州电厂", "金州", "金州厂"]),
        ("北方热电", ["北方热电", "北方电厂", "北方", "北方厂"]),
        ("庄河环海", ["庄河环海", "庄河电厂", "庄河", "庄河厂"]),
        ("金普公司", ["金普公司", "金普热电", "金普", "金普新区"]),
        ("供热公司", ["供热公司", "管网公司", "供热", "管网"]),
        ("大连港", ["大连港"]),
        ("调度运行中心", ["调度中心", "运行中心", "调度运行中心"])
    ]

    CATEGORY_ORG_MAP = {
        "北海": "北海热电联产",
        "水炉": "北海水炉",
        "燃气炉": "燃气锅炉",
        "香海": "香海电厂",
        "金州": "金州热电",
        "北方": "北方热电",
        "庄河": "庄河环海"
    }

    def __init__(self, paragraphs: List[str], tables: List[List[List[str]]], file_name: str = "", project_key: str = "daily_report_25_26"):
        self.paragraphs = paragraphs
        self.tables = tables
        self.file_name = file_name
        self.project_key = project_key

        # 智能表格语义分类插槽
        self.table_boiler_turbine: Optional[List[List[str]]] = None
        self.table_first_stations: Optional[List[List[str]]] = None
        self.table_steam_networks: Optional[List[List[str]]] = None
        self.table_coal_stats: Optional[List[List[str]]] = None
        self.table_production: Optional[List[List[str]]] = None
        self.table_network_heat: Optional[List[List[str]]] = None
        self.table_event_logs: Optional[List[List[str]]] = None

        self._classify_tables()

    @staticmethod
    def clean_text(text: Optional[str]) -> str:
        """文本清洗：统一全半角、异体字替换、去除多余空白"""
        if not text:
            return ""
        t = str(text).replace("幷", "并").replace("℃", "℃")
        t = t.replace("～", "~").replace("—", "-").replace("－", "-")
        t = t.replace("：", ":").replace("，", ",").replace("／", "/")
        t = re.sub(r"\s+", " ", t).strip()
        return t

    @staticmethod
    def parse_float(val: Any) -> Optional[float]:
        """安全提取浮点数"""
        if val is None:
            return None
        s = str(val).strip()
        if not s or s in ["-", "--", "/", "null", "None", "无"]:
            return None
        match = re.search(r"[-+]?\d*\.?\d+", s)
        if match:
            try:
                return float(match.group())
            except ValueError:
                return None
        return None

    @staticmethod
    def parse_int(val: Any) -> Optional[int]:
        """安全提取整数"""
        f = UnifiedDispatchExtractor.parse_float(val)
        return int(f) if f is not None else None

    @staticmethod
    def expand_device_numbers(text: str) -> List[str]:
        """展开设备编号：如 #1-10 或 #2、3"""
        cleaned = UnifiedDispatchExtractor.clean_text(text)
        if not cleaned:
            return []

        results = []
        range_match = re.search(r"#?(\d+)\s*[-~]\s*#?(\d+)", cleaned)
        if range_match:
            start_num = int(range_match.group(1))
            end_num = int(range_match.group(2))
            for n in range(start_num, end_num + 1):
                results.append(f"#{n}")
            return results

        nums = re.findall(r"\d+", cleaned)
        for n in nums:
            results.append(f"#{n}")
        return sorted(list(set(results)), key=lambda x: int(x.replace("#", "")))

    @classmethod
    def standardize_boiler_name(cls, raw_text: str, org_or_unit: str) -> Tuple[str, str, str]:
        """
        根据集团底层拓扑规则精准推导锅炉类型与标准命名：
        返回: (boiler_type, standard_facility_name, sub_unit)
        """
        t = raw_text.strip()
        num_m = re.search(r"(\d+)", t)
        num = f"#{num_m.group(1)}" if num_m else ""

        # 庄河环海：全为锅炉房锅炉
        if any(k in org_or_unit for k in ["庄河", "庄河环海", "庄河电厂"]):
            return "锅炉房锅炉", f"庄河{num}锅炉房锅炉" if num else "庄河锅炉房锅炉", "庄河环海"

        # 金州热电：#6炉是水炉，其余全为汽炉
        if any(k in org_or_unit for k in ["金州", "金州热电", "金州电厂"]):
            if num == "#6" or "6" in t or "水炉" in t:
                return "水炉", "#6水炉", "金州热电"
            else:
                return "汽炉", f"{num}汽炉" if num else "金州汽炉", "金州热电"

        # 北方热电：全为汽炉
        if any(k in org_or_unit for k in ["北方", "北方热电", "北方电厂"]):
            return "汽炉", f"{num}汽炉" if num else "北方汽炉", "北方热电"

        # 香海电厂：全为汽炉
        if any(k in org_or_unit for k in ["香海", "香海电厂", "香海厂"]):
            return "汽炉", f"{num}汽炉" if num else "香海汽炉", "香海电厂"

        # 北海电厂：细分热电联产(汽炉)、水炉(水炉)、燃气炉(燃气炉)
        if "水炉" in t or "水炉" in org_or_unit:
            return "水炉", f"{num}水炉" if num else "北海水炉", "北海水炉"
        elif "燃气" in t or "燃气" in org_or_unit:
            return "燃气炉", f"{num}燃气炉" if num else "北海燃气炉", "燃气锅炉"
        else:
            return "汽炉", f"{num}汽炉" if num else "北海汽炉", "北海热电联产"

    @classmethod
    def standardize_org_name(cls, text: str) -> str:
        """将任意提及的企业名称映射至标准字典名称"""
        if not text:
            return "调度运行中心"
        t = text.strip()
        for std_name, aliases in cls.KNOWN_ORGS:
            for alias in aliases:
                if alias in t:
                    return std_name
        return t

    def _classify_tables(self):
        """核心智能分类器：根据表格表头特征语义自适应识别各个表格"""
        for tbl in self.tables:
            if not tbl or len(tbl) == 0:
                continue

            feature_text = " ".join([self.clean_text(c) for row in tbl[:6] for c in row])
            feature_all = " ".join([self.clean_text(c) for row in tbl for c in row])

            # 1. 机炉运行情况表
            if ("运行" in feature_text and "备用" in feature_text and "检修" in feature_text) and \
               ("分类" in feature_text or "北海" in feature_text or "机" in feature_text or "炉" in feature_text):
                if not self.table_boiler_turbine:
                    self.table_boiler_turbine = tbl
                    continue

            # 2. 水网首站工况表
            if ("水网首站" in feature_text or "首站" in feature_text) and \
               ("循环量" in feature_text or "供回压力" in feature_text or "供回温" in feature_text or "运行方式" in feature_text):
                if not self.table_first_stations:
                    self.table_first_stations = tbl
                    continue

            # 3. 蒸汽网口情况表
            if ("蒸汽网口" in feature_text or "网口" in feature_text) and \
               ("运行状态" in feature_text or "压力" in feature_text or "流量" in feature_text):
                if not self.table_steam_networks:
                    self.table_steam_networks = tbl
                    continue

            # 4. 燃煤统计表
            if "进煤" in feature_text and "耗煤" in feature_text and "存煤" in feature_text:
                if not self.table_coal_stats:
                    self.table_coal_stats = tbl
                    continue

            # 5. 当日产量统计表
            if ("当日计划" in feature_text and "当日实际" in feature_text) or \
               ("汽炉耗煤" in feature_all or "水炉耗煤" in feature_all or "耗煤总量" in feature_all or "当日产量" in feature_text):
                if not self.table_production:
                    self.table_production = tbl
                    continue

            # 6. 网口热量表 (独立分表存在时)
            if "网口热量" in feature_text or ("高温水" in feature_text and "低真空" in feature_text):
                if not self.table_network_heat:
                    self.table_network_heat = tbl
                    continue

            # 7. 生产信息记录与微观事件表
            if "生产运行情况" in feature_text or "生产信息记录" in feature_text or \
               ("时间" in feature_text and ("运行" in feature_text or "情况" in feature_text)):
                if not self.table_event_logs:
                    self.table_event_logs = tbl
                    continue

    @staticmethod
    def safe_format_date(y: int, m: int, d: int) -> Optional[str]:
        """安全格式化日期，自动钳位处理平年2月29日等手误"""
        try:
            dt = datetime(y, m, d)
            return dt.strftime("%Y-%m-%d")
        except ValueError:
            if m == 2 and d >= 29:
                return f"{y:04d}-02-28"
            elif d > 30 and m in [4, 6, 9, 11]:
                return f"{y:04d}-{m:02d}-30"
            elif d > 31:
                return f"{y:04d}-{m:02d}-31"
            return None

    def extract_header(self) -> Dict[str, Any]:
        """提取值班人员全员、业务日期与发布日期"""
        duty_officers = []
        publish_date = None
        biz_date = None
        issue_date = None

        full_doc_text = " ".join(self.paragraphs)
        if self.tables:
            for row in self.tables[0]:
                full_doc_text += " " + " ".join(row)

        duty_match = re.search(r"值班人员\s*[:：]\s*(.*?)(?:时间|$)", full_doc_text)
        if duty_match:
            raw_officers = duty_match.group(1).strip()
            officers_list = [o.strip() for o in re.split(r"[\s,、，]+", raw_officers) if o.strip()]
            duty_officers = officers_list

        time_match = re.search(r"时间\s*[:：]\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", full_doc_text)
        if time_match:
            y, m, d = int(time_match.group(1)), int(time_match.group(2)), int(time_match.group(3))
            biz_date = self.safe_format_date(y, m, d)

        pub_match = re.search(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日\s*$", full_doc_text.strip())
        if not pub_match:
            for p in reversed(self.paragraphs[-6:]):
                p_c = self.clean_text(p)
                m = re.search(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", p_c)
                if m:
                    y, m_val, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
                    publish_date = self.safe_format_date(y, m_val, d)
                    issue_date = publish_date
                    break
        else:
            y, m_val, d = int(pub_match.group(1)), int(pub_match.group(2)), int(pub_match.group(3))
            publish_date = self.safe_format_date(y, m_val, d)
            issue_date = publish_date

        if not biz_date and self.file_name:
            fn_match = re.search(r"(\d{4})年?(\d{1,2})月(\d{1,2})日", self.file_name)
            if fn_match:
                y, m, d = int(fn_match.group(1)), int(fn_match.group(2)), int(fn_match.group(3))
                biz_date = self.safe_format_date(y, m, d)

        return {
            "duty_officers": duty_officers,
            "biz_date": biz_date,
            "issue_date": issue_date or publish_date,
            "publish_date": publish_date
        }

    def extract_weather(self, biz_date: Optional[str] = None) -> Dict[str, Any]:
        """提取气象环境全要素物理指标（支持多行跨行合并与全角波浪号清洗）"""
        weather_lines = []
        capturing = False

        for p in self.paragraphs:
            p_clean = self.clean_text(p)
            if not p_clean:
                continue

            if "天气情况" in p_clean:
                capturing = True
                after_title = re.sub(r"^[一二三四五六七八九0-9\s、.．]*天气情况[:：]?", "", p_clean).strip()
                if after_title:
                    weather_lines.append(after_title)
                continue

            if capturing:
                # 遇到下一个大章节小标题（如 "二、电厂主设备运行情况"）终止
                if re.search(r"^[一二三四五六七八九0-9\s、.．]+[、.．]?\s*(电厂|主设备|生产|首站)", p_clean) or \
                   "主设备运行情况" in p_clean or "生产运行情况" in p_clean:
                    break
                weather_lines.append(p_clean)
                if len(weather_lines) >= 3:
                    break

        if not weather_lines:
            for p in self.paragraphs:
                p_clean = self.clean_text(p)
                if ("℃" in p_clean or "相对湿度" in p_clean) and ("风" in p_clean or "晴" in p_clean or "雨" in p_clean or "阴" in p_clean):
                    weather_lines.append(p_clean)
                    break

        # PDF 把“相对湿度”和风级拆到多行；气象字段内空白无业务含义。
        weather_raw_text = " ".join(weather_lines).strip()
        weather_text = re.sub(r'\s+', '', weather_raw_text)

        res = {
            "raw_text": weather_raw_text,
            "forecast_date": None,
            "weekday": "",
            "weather_desc": "",
            "temp_min": None,
            "temp_max": None,
            "humidity_min": None,
            "humidity_max": None,
            "wind": "",
            "wind_direction": "",
            "wind_scale": "",
            "humidity": ""
        }
        if not weather_text:
            return res

        wth_date_m = re.search(r"(\d{1,2})\s*月\s*(\d{1,2})\s*日", weather_text)
        if wth_date_m and biz_date:
            wm, wd = wth_date_m.groups()
            year = int(biz_date.split("-")[0])
            res["forecast_date"] = self.safe_format_date(year, int(wm), int(wd))

        wk_m = re.search(r"(星期[一二三四五六日天])", weather_text)
        if wk_m:
            res["weekday"] = wk_m.group(1)

        temp_m = re.search(r"(-?\d+(?:\.\d+)?)\s*℃\s*[-~至到～]\s*(-?\d+(?:\.\d+)?)\s*℃", weather_text)
        if temp_m:
            t1 = float(temp_m.group(1))
            t2 = float(temp_m.group(2))
            res["temp_min"] = min(t1, t2)
            res["temp_max"] = max(t1, t2)
        else:
            single_temp_m = re.search(r"(-?\d+(?:\.\d+)?)\s*℃", weather_text)
            if single_temp_m:
                res["temp_min"] = float(single_temp_m.group(1))
                res["temp_max"] = float(single_temp_m.group(1))

        hum_m = re.search(r"相对湿度\s*[:：]?\s*(\d+)\s*%\s*[-~至到～]\s*(\d+)\s*%", weather_text)
        if hum_m:
            h1 = int(hum_m.group(1))
            h2 = int(hum_m.group(2))
            res["humidity_min"] = min(h1, h2)
            res["humidity_max"] = max(h1, h2)
            res["humidity"] = f"{min(h1, h2)}%~{max(h1, h2)}%"

        desc_m = re.search(r"℃([^\d℃~]+?)(?=相对湿度|(?:东|西|南|北|偏|微)+风|$)", weather_text)
        if desc_m:
            res["weather_desc"] = desc_m.group(1).strip()

        wind_m = re.search(r"((?:东|西|南|北|偏|微)+风)([\d~～级阵风转偏东西南北微\-]+)", weather_text)
        if wind_m:
            clean_dir = re.sub(r"[\s\u00a0]+", "", wind_m.group(1))
            res["wind_direction"] = clean_dir
            res["wind_scale"] = wind_m.group(2).strip()
            res["wind"] = f"{res['wind_direction']}{res['wind_scale']}"

        return res

    def extract_boiler_turbine_status(self) -> Dict[str, Any]:
        """提取机炉运行情况并执行拓扑深度归一"""
        if not self.table_boiler_turbine:
            return {"facilities": [], "matrix": {}, "flat_map": {}}

        table = self.table_boiler_turbine
        facilities = []
        matrix: Dict[str, Any] = {}
        flat_map: Dict[str, str] = {}

        headers = [self.clean_text(c) for c in table[0]]
        status_cols = {}
        for col_idx, h in enumerate(headers):
            if "运行" in h:
                status_cols["运行"] = col_idx
            elif "备用" in h:
                status_cols["备用"] = col_idx
            elif "检修" in h:
                status_cols["检修"] = col_idx

        if not status_cols:
            status_cols = {"运行": 1, "备用": 2, "检修": 3}

        for row_idx, row in enumerate(table):
            if row_idx == 0:
                continue
            cells = [self.clean_text(c) for c in row]
            if not any(cells) or not cells[0]:
                continue

            raw_cat = cells[0].replace(" ", "")
            std_unit = self.CATEGORY_ORG_MAP.get(raw_cat, raw_cat)

            if std_unit not in matrix:
                matrix[std_unit] = {
                    "机": {"运行": [], "备用": [], "检修": []},
                    "炉": {"运行": [], "备用": [], "检修": []}
                }

            for status_name, col_idx in status_cols.items():
                if col_idx >= len(cells):
                    continue
                content = cells[col_idx]
                if not content or content in ["-", "--", "/", "无"]:
                    continue

                m_part = ""
                b_part = ""
                if "机" in content and "炉" in content:
                    parts = content.split("机")
                    m_part = parts[0]
                    b_part = parts[1]
                elif "机" in content:
                    m_part = content
                elif "炉" in content:
                    b_part = content
                else:
                    if raw_cat in ["水炉", "燃气炉"]:
                        b_part = content
                    else:
                        m_part = content

                m_list = self.expand_device_numbers(m_part)
                b_list = self.expand_device_numbers(b_part)

                matrix[std_unit]["机"][status_name].extend(m_list)
                matrix[std_unit]["炉"][status_name].extend(b_list)

                for dev in m_list:
                    dev_key = f"{std_unit}_{dev}机"
                    flat_map[dev_key] = status_name
                    facilities.append({
                        "facility_type": "汽机",
                        "raw_name": f"{dev}机",
                        "standard_name": f"{dev}汽机",
                        "unit_or_plant": std_unit,
                        "status": "启动中" if re.search(
                            rf"(?<!\d)#?{dev.lstrip('#')}\s*(?:汽机|机组|机)?\s*[（(][^）)]*启动中", content
                        ) else status_name,
                        "table_text": content,
                    })

                for dev in b_list:
                    b_type, b_std_name, sub_u = self.standardize_boiler_name(dev, raw_cat)
                    dev_key = f"{std_unit}_{b_std_name}"
                    flat_map[dev_key] = status_name
                    facilities.append({
                        "facility_type": "锅炉",
                        "boiler_type": b_type,
                        "raw_name": f"{dev}炉",
                        "standard_name": b_std_name,
                        "unit_or_plant": sub_u,
                        "status": "启动中" if re.search(
                            rf"(?<!\d)#?{dev.lstrip('#')}\s*(?:汽炉|水炉|燃气炉|炉)?\s*[（(][^）)]*启动中", content
                        ) else status_name,
                        "table_text": content,
                    })

        return {
            "facilities": facilities,
            "matrix": matrix,
            "flat_map": flat_map
        }

    def extract_first_stations(self) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
        """
        提取水网首站运行工况
        返回: (stations_list, flat_stations_metrics_36)
        """
        flat_metrics = {
            "bh_s1_pumps": None, "bh_s1_flow": None, "bh_s1_press_sup": None, "bh_s1_press_ret": None, "bh_s1_temp_sup": None, "bh_s1_temp_ret": None,
            "bh_s2_pumps": None, "bh_s2_flow": None, "bh_s2_press_sup": None, "bh_s2_press_ret": None, "bh_s2_temp_sup": None, "bh_s2_temp_ret": None,
            "bh_wl_pumps": None, "bh_wl_flow": None, "bh_wl_press_sup": None, "bh_wl_press_ret": None, "bh_wl_temp_sup": None, "bh_wl_temp_ret": None,
            "xh_lv1_pumps": None, "xh_lv1_flow": None, "xh_lv1_press_sup": None, "xh_lv1_press_ret": None, "xh_lv1_temp_sup": None, "xh_lv1_temp_ret": None,
            "xh_lv3_pumps": None, "xh_lv3_flow": None, "xh_lv3_press_sup": None, "xh_lv3_press_ret": None, "xh_lv3_temp_sup": None, "xh_lv3_temp_ret": None,
            "xh_main_pumps": None, "xh_main_flow": None, "xh_main_press_sup": None, "xh_main_press_ret": None, "xh_main_temp_sup": None, "xh_main_temp_ret": None,
        }

        if not self.table_first_stations:
            return [], flat_metrics

        table = self.table_first_stations
        stations = []
        current_plant = "北海电厂"

        for row_idx, row in enumerate(table):
            if row_idx == 0:
                continue
            cells = [self.clean_text(c) for c in row]
            if not any(cells):
                continue

            plant_cell = ""
            station_cell = ""
            mode_cell = ""
            flow_cell = ""
            press_cell = ""
            temp_cell = ""

            if len(cells) >= 6:
                plant_cell = cells[0]
                station_cell = cells[1]
                mode_cell = cells[2]
                flow_cell = cells[3]
                press_cell = cells[4]
                temp_cell = cells[5]
            elif len(cells) == 5:
                station_cell = cells[0]
                mode_cell = cells[1]
                flow_cell = cells[2]
                press_cell = cells[3]
                temp_cell = cells[4]
            else:
                continue

            if plant_cell:
                if "北海" in plant_cell:
                    current_plant = "北海电厂"
                elif "香海" in plant_cell:
                    current_plant = "香海电厂"
                else:
                    current_plant = self.standardize_org_name(plant_cell)

            if not station_cell:
                continue

            p_in, p_out = None, None
            if press_cell and "/" in press_cell:
                p_parts = press_cell.split("/")
                p_in = self.parse_float(p_parts[0])
                p_out = self.parse_float(p_parts[1]) if len(p_parts) > 1 else None

            t_in, t_out = None, None
            if temp_cell and "/" in temp_cell:
                t_parts = temp_cell.split("/")
                t_in = self.parse_float(t_parts[0])
                t_out = self.parse_float(t_parts[1]) if len(t_parts) > 1 else None

            pump_m = re.search(r"(\d+)\s*台", mode_cell)
            pump_count = int(pump_m.group(1)) if pump_m else self.parse_int(mode_cell)
            flow_val = self.parse_float(flow_cell)

            stations.append({
                "plant": current_plant,
                "station_name": station_cell,
                "pump_count": pump_count,
                "flow_rate_th": flow_val,
                "supply_pressure_mpa": p_in,
                "return_pressure_mpa": p_out,
                "pressure_diff_mpa": round(p_in - p_out, 3) if p_in is not None and p_out is not None else None,
                "supply_temp_c": t_in,
                "return_temp_c": t_out,
                "temp_diff_c": round(t_in - t_out, 2) if t_in is not None and t_out is not None else None
            })

            # 填充到 36 项物理大宽表字段
            st_key = None
            if "#1首站" in station_cell or "#1 首站" in station_cell:
                st_key = "bh_s1"
            elif "#2首站" in station_cell or "#2 首站" in station_cell:
                st_key = "bh_s2"
            elif "水炉" in station_cell and current_plant == "北海电厂":
                st_key = "bh_wl"
            elif "#1低真空" in station_cell or "#1 低真空" in station_cell:
                st_key = "xh_lv1"
            elif "#3低真空" in station_cell or "#3 低真空" in station_cell:
                st_key = "xh_lv3"
            elif "首站" in station_cell and current_plant == "香海电厂":
                st_key = "xh_main"

            if st_key:
                flat_metrics[f"{st_key}_pumps"] = pump_count
                flat_metrics[f"{st_key}_flow"] = flow_val
                flat_metrics[f"{st_key}_press_sup"] = p_in
                flat_metrics[f"{st_key}_press_ret"] = p_out
                flat_metrics[f"{st_key}_temp_sup"] = t_in
                flat_metrics[f"{st_key}_temp_ret"] = t_out

        return stations, flat_metrics

    def extract_steam_networks(self) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
        """
        提取蒸汽网口工况
        返回: (networks_list, flat_steam_metrics_12)
        """
        flat_metrics = {
            "bh_south_status": None, "bh_south_press": None, "bh_south_flow": None,
            "xb_link_status": None,  "xb_link_press": None,  "xb_link_flow": None,
            "xh_east_status": None,  "xh_east_press": None,  "xh_east_flow": None,
            "xh_west_status": None,  "xh_west_press": None,  "xh_west_flow": None,
        }

        if not self.table_steam_networks:
            return [], flat_metrics

        table = self.table_steam_networks
        networks = []

        for row_idx, row in enumerate(table):
            if row_idx == 0:
                continue
            cells = [self.clean_text(c) for c in row]
            if not any(cells) or len(cells) < 2:
                continue

            name = cells[0].replace(" ", "")
            status = cells[1] if len(cells) > 1 else "投运"
            press_raw = cells[2] if len(cells) > 2 else ""
            flow_raw = cells[3] if len(cells) > 3 else ""

            if "/" in press_raw or "/" in flow_raw:
                p_parts = press_raw.split("/") if "/" in press_raw else [press_raw, press_raw]
                f_parts = flow_raw.split("/") if "/" in flow_raw else [flow_raw, flow_raw]
                networks.append({
                    "network_name": "香海东网",
                    "status": status,
                    "pressure_mpa": self.parse_float(p_parts[0]),
                    "flow_th": self.parse_float(f_parts[0])
                })
                networks.append({
                    "network_name": "香海西网",
                    "status": status,
                    "pressure_mpa": self.parse_float(p_parts[1]) if len(p_parts) > 1 else None,
                    "flow_th": self.parse_float(f_parts[1]) if len(f_parts) > 1 else None
                })
                flat_metrics["xh_east_status"] = status
                flat_metrics["xh_east_press"] = self.parse_float(p_parts[0])
                flat_metrics["xh_east_flow"] = self.parse_float(f_parts[0])
                flat_metrics["xh_west_status"] = status
                flat_metrics["xh_west_press"] = self.parse_float(p_parts[1]) if len(p_parts) > 1 else None
                flat_metrics["xh_west_flow"] = self.parse_float(f_parts[1]) if len(f_parts) > 1 else None
            else:
                p_val = self.parse_float(press_raw)
                f_val = self.parse_float(flow_raw)
                networks.append({
                    "network_name": name,
                    "status": status,
                    "pressure_mpa": p_val,
                    "flow_th": f_val
                })
                if "北海南网" in name:
                    flat_metrics["bh_south_status"] = status
                    flat_metrics["bh_south_press"] = p_val
                    flat_metrics["bh_south_flow"] = f_val
                elif "香北联络" in name:
                    flat_metrics["xb_link_status"] = status
                    flat_metrics["xb_link_press"] = p_val
                    flat_metrics["xb_link_flow"] = f_val
                elif "香海东" in name:
                    flat_metrics["xh_east_status"] = status
                    flat_metrics["xh_east_press"] = p_val
                    flat_metrics["xh_east_flow"] = f_val
                elif "香海西" in name:
                    flat_metrics["xh_west_status"] = status
                    flat_metrics["xh_west_press"] = p_val
                    flat_metrics["xh_west_flow"] = f_val

        return networks, flat_metrics

    def extract_coal_stats(self) -> Tuple[Dict[str, Any], Dict[str, Any]]:
        """
        提取燃煤收耗存统计表与进煤说明
        返回: (coal_dict, flat_coal_metrics_25)
        """
        flat_metrics = {
            "coal_remark": "",
            "coal_bh_in": None, "coal_bh_use": None, "coal_bh_stock": None,
            "coal_xh_in": None, "coal_xh_use": None, "coal_xh_stock": None,
            "coal_main_sum_in": None, "coal_main_sum_use": None, "coal_main_sum_stock": None,
            "coal_jz_in": None, "coal_jz_use": None, "coal_jz_stock": None,
            "coal_bf_in": None, "coal_bf_use": None, "coal_bf_stock": None,
            "coal_zh_in": None, "coal_zh_use": None, "coal_zh_stock": None,
            "coal_dlg_in": None, "coal_dlg_use": None, "coal_dlg_stock": None,
            "coal_total_in": None, "coal_total_use": None, "coal_total_stock": None,
        }

        if not self.table_coal_stats:
            return {"coal_items": [], "coal_remark": ""}, flat_metrics

        table = self.table_coal_stats
        items = []
        remark_lines = []

        for row_idx, row in enumerate(table):
            if row_idx == 0:
                continue
            cells = [self.clean_text(c) for c in row]
            if not any(cells):
                continue

            first_col = cells[0].replace(" ", "")

            if "进煤:" in first_col or "进煤：" in first_col or "汽运车" in first_col:
                remark_lines.append(cells[0].strip())
                continue

            if len(cells) < 4:
                continue

            cat = cells[0]
            in_coal = self.parse_float(cells[1])
            consumed_coal = self.parse_float(cells[2])
            storage_coal = self.parse_float(cells[3])

            org_name = cat
            entity_key = None
            if "北海" in cat:
                org_name = "北海电厂"
                entity_key = "bh"
            elif "香海" in cat:
                org_name = "香海电厂"
                entity_key = "xh"
            elif "主城区" in cat:
                org_name = "主城区合计"
                entity_key = "main_sum"
            elif "金州" in cat:
                org_name = "金州热电"
                entity_key = "jz"
            elif "北方" in cat:
                org_name = "北方热电"
                entity_key = "bf"
            elif "庄河" in cat:
                org_name = "庄河环海"
                entity_key = "zh"
            elif "金普" in cat:
                org_name = "金普公司"
            elif "大连港" in cat:
                org_name = "大连港"
                entity_key = "dlg"
            elif "合计" in cat:
                org_name = "全网合计"
                entity_key = "total"

            items.append({
                "category": cat,
                "standard_org": org_name,
                "in_coal_tons": in_coal,
                "consumed_coal_tons": consumed_coal,
                "storage_coal_tons": storage_coal
            })

            if entity_key:
                flat_metrics[f"coal_{entity_key}_in"] = in_coal
                flat_metrics[f"coal_{entity_key}_use"] = consumed_coal
                flat_metrics[f"coal_{entity_key}_stock"] = storage_coal

        remark_str = "；".join(list(set(remark_lines)))
        flat_metrics["coal_remark"] = remark_str

        return {
            "coal_items": items,
            "coal_remark": remark_str
        }, flat_metrics

    def extract_production_and_network_heat(self) -> Tuple[Dict[str, Any], Dict[str, Any], Dict[str, Any]]:
        """
        提取当日产量统计与网口热量（自适应支持合并大表与独立分表）
        返回: (production_stats, network_heat, flat_prod_gw_metrics_26)
        """
        prod_res = {
            "beihai_plan_coal": None,
            "beihai_actual_coal": None,
            "beihai_plan_power": None,
            "beihai_actual_power": None,
            "beihai_plan_heat": None,
            "beihai_actual_heat": None,
            "xianghai_plan_coal": None,
            "xianghai_actual_coal": None,
            "xianghai_plan_power": None,
            "xianghai_actual_power": None,
            "xianghai_plan_heat": None,
            "xianghai_actual_heat": None,
            "urban_total_plan_coal": None,
            "urban_total_actual_coal": None,
            "urban_total_plan_power": None,
            "urban_total_actual_power": None,
            "urban_total_plan_heat": None,
            "urban_total_actual_heat": None,
            "beihai_steam_coal_plan": None,
            "beihai_steam_coal_act": None,
            "beihai_water_coal_plan": None,
            "beihai_water_coal_act": None
        }

        heat_res = {
            "gw_high_temp_water": None,
            "gw_lv1_heat": None,
            "gw_lv3_heat": None,
            "gw_bh_south_steam": None,
            "gw_xh_ew_steam": None,
            "gw_xb_link_steam": None
        }

        flat_metrics = {
            "prod_bh_steam_coal_plan": None, "prod_bh_steam_coal_act": None,
            "prod_bh_water_coal_plan": None, "prod_bh_water_coal_act": None,
            "prod_bh_power_plan": None,      "prod_bh_power_act": None,
            "prod_bh_heat_plan": None,       "prod_bh_heat_act": None,
            "prod_xh_steam_coal_plan": None, "prod_xh_steam_coal_act": None,
            "prod_xh_power_plan": None,      "prod_xh_power_act": None,
            "prod_xh_heat_plan": None,       "prod_xh_heat_act": None,
            "prod_main_coal_plan": None,     "prod_main_coal_act": None,
            "prod_main_power_plan": None,    "prod_main_power_act": None,
            "prod_main_heat_plan": None,     "prod_main_heat_act": None,
            "gw_high_temp_water": None,
            "gw_lv1_heat": None,
            "gw_lv3_heat": None,
            "gw_bh_south_steam": None,
            "gw_xh_ew_steam": None,
            "gw_xb_link_steam": None
        }

        tables_to_scan = []
        if self.table_production:
            tables_to_scan.append(self.table_production)
        if self.table_network_heat:
            tables_to_scan.append(self.table_network_heat)

        for tbl in tables_to_scan:
            for row in tbl:
                cells = [self.clean_text(c) for c in row]
                if not any(cells):
                    continue

                first_col = cells[0].replace(" ", "")

                if "汽炉耗煤" in first_col:
                    if len(cells) > 2:
                        p_b = self.parse_float(cells[1])
                        a_b = self.parse_float(cells[2])
                        prod_res["beihai_steam_coal_plan"] = p_b
                        prod_res["beihai_steam_coal_act"] = a_b
                        flat_metrics["prod_bh_steam_coal_plan"] = p_b
                        flat_metrics["prod_bh_steam_coal_act"] = a_b
                    if len(cells) > 4:
                        p_x = self.parse_float(cells[3])
                        a_x = self.parse_float(cells[4])
                        prod_res["xianghai_plan_coal"] = p_x
                        prod_res["xianghai_actual_coal"] = a_x
                        flat_metrics["prod_xh_steam_coal_plan"] = p_x
                        flat_metrics["prod_xh_steam_coal_act"] = a_x

                elif "水炉耗煤" in first_col:
                    if len(cells) > 2:
                        p_bw = self.parse_float(cells[1])
                        a_bw = self.parse_float(cells[2])
                        prod_res["beihai_water_coal_plan"] = p_bw
                        prod_res["beihai_water_coal_act"] = a_bw
                        flat_metrics["prod_bh_water_coal_plan"] = p_bw
                        flat_metrics["prod_bh_water_coal_act"] = a_bw

                elif "发电" in first_col and "总量" not in first_col:
                    if len(cells) > 2:
                        p_bp = self.parse_float(cells[1])
                        a_bp = self.parse_float(cells[2])
                        prod_res["beihai_plan_power"] = p_bp
                        prod_res["beihai_actual_power"] = a_bp
                        flat_metrics["prod_bh_power_plan"] = p_bp
                        flat_metrics["prod_bh_power_act"] = a_bp
                    if len(cells) > 4:
                        p_xp = self.parse_float(cells[3])
                        a_xp = self.parse_float(cells[4])
                        prod_res["xianghai_plan_power"] = p_xp
                        prod_res["xianghai_actual_power"] = a_xp
                        flat_metrics["prod_xh_power_plan"] = p_xp
                        flat_metrics["prod_xh_power_act"] = a_xp

                elif "供热量" in first_col and "总量" not in first_col:
                    if len(cells) > 2:
                        p_bh = self.parse_float(cells[1])
                        a_bh = self.parse_float(cells[2])
                        prod_res["beihai_plan_heat"] = p_bh
                        prod_res["beihai_actual_heat"] = a_bh
                        flat_metrics["prod_bh_heat_plan"] = p_bh
                        flat_metrics["prod_bh_heat_act"] = a_bh
                    if len(cells) > 4:
                        p_xh = self.parse_float(cells[3])
                        a_xh = self.parse_float(cells[4])
                        prod_res["xianghai_plan_heat"] = p_xh
                        prod_res["xianghai_actual_heat"] = a_xh
                        flat_metrics["prod_xh_heat_plan"] = p_xh
                        flat_metrics["prod_xh_heat_act"] = a_xh

                elif "耗煤总量" in first_col:
                    p_c = self.parse_float(cells[1]) if len(cells) > 1 else None
                    a_c = self.parse_float(cells[3]) if len(cells) > 3 else self.parse_float(cells[2]) if len(cells) > 2 else None
                    prod_res["urban_total_plan_coal"] = p_c
                    prod_res["urban_total_actual_coal"] = a_c
                    flat_metrics["prod_main_coal_plan"] = p_c
                    flat_metrics["prod_main_coal_act"] = a_c

                elif "发电总量" in first_col:
                    p_pw = self.parse_float(cells[1]) if len(cells) > 1 else None
                    a_pw = self.parse_float(cells[3]) if len(cells) > 3 else self.parse_float(cells[2]) if len(cells) > 2 else None
                    prod_res["urban_total_plan_power"] = p_pw
                    prod_res["urban_total_actual_power"] = a_pw
                    flat_metrics["prod_main_power_plan"] = p_pw
                    flat_metrics["prod_main_power_act"] = a_pw

                elif "供热总量" in first_col:
                    p_ht = self.parse_float(cells[1]) if len(cells) > 1 else None
                    a_ht = self.parse_float(cells[3]) if len(cells) > 3 else self.parse_float(cells[2]) if len(cells) > 2 else None
                    prod_res["urban_total_plan_heat"] = p_ht
                    prod_res["urban_total_actual_heat"] = a_ht
                    flat_metrics["prod_main_heat_plan"] = p_ht
                    flat_metrics["prod_main_heat_act"] = a_ht

                elif "高温水" in first_col:
                    v = self.parse_float(cells[-1])
                    heat_res["gw_high_temp_water"] = v
                    flat_metrics["gw_high_temp_water"] = v
                elif "#1低真空" in first_col or "#1 低真空" in first_col:
                    v = self.parse_float(cells[-1])
                    heat_res["gw_lv1_heat"] = v
                    flat_metrics["gw_lv1_heat"] = v
                elif "#3低真空" in first_col or "#3 低真空" in first_col:
                    v = self.parse_float(cells[-1])
                    heat_res["gw_lv3_heat"] = v
                    flat_metrics["gw_lv3_heat"] = v
                elif "北海南网供汽量" in first_col:
                    v = self.parse_float(cells[-1])
                    heat_res["gw_bh_south_steam"] = v
                    flat_metrics["gw_bh_south_steam"] = v
                elif "香海东西网供汽量" in first_col or "香海东、西网" in first_col:
                    v = self.parse_float(cells[-1])
                    heat_res["gw_xh_ew_steam"] = v
                    flat_metrics["gw_xh_ew_steam"] = v
                elif "香北联络" in first_col:
                    v = self.parse_float(cells[-1])
                    heat_res["gw_xb_link_steam"] = v
                    flat_metrics["gw_xb_link_steam"] = v

        bh_c_act = 0.0
        has_bh = False
        if prod_res["beihai_steam_coal_act"] is not None:
            bh_c_act += prod_res["beihai_steam_coal_act"]
            has_bh = True
        if prod_res["beihai_water_coal_act"] is not None:
            bh_c_act += prod_res["beihai_water_coal_act"]
            has_bh = True
        prod_res["beihai_actual_coal"] = bh_c_act if has_bh else None

        bh_c_plan = 0.0
        has_bh_p = False
        if prod_res["beihai_steam_coal_plan"] is not None:
            bh_c_plan += prod_res["beihai_steam_coal_plan"]
            has_bh_p = True
        if prod_res["beihai_water_coal_plan"] is not None:
            bh_c_plan += prod_res["beihai_water_coal_plan"]
            has_bh_p = True
        prod_res["beihai_plan_coal"] = bh_c_plan if has_bh_p else None

        return prod_res, heat_res, flat_metrics

    @classmethod
    def detect_org_header(cls, time_col: str, content_col: str) -> Optional[str]:
        """组织分组头识别"""
        c = content_col.strip()
        t = time_col.strip()
        c_clean = c.replace(" ", "").rstrip(":：")
        t_clean = t.replace(" ", "").rstrip(":：")

        if len(c_clean) > 12:
            return None

        for std_name, aliases in cls.KNOWN_ORGS:
            for a in aliases:
                if (c_clean == a or t_clean == a) and not t_clean.replace(":", "").isdigit():
                    return std_name
                if c.endswith((":", "：")) and (c_clean == a or c_clean.startswith(a) or c_clean.endswith(a)):
                    return std_name
        return None

    @staticmethod
    def split_compound_event(time_str: str, desc: str) -> List[Tuple[str, str]]:
        """
        智能拆分包含多个明确时刻的复合操作事件（如: 10:10 投运... 22:17 跳闸倒泵）
        注意：工作票工期严禁拆分
        """
        if "工作票" in desc or "工期" in desc:
            return [(time_str, desc)]

        pattern = re.compile(r"([。；;\n!！,，]|\s+)(\d{1,2}:\d{2})")
        matches = list(pattern.finditer(desc))
        if not matches:
            return [(time_str, desc)]

        sub_events = []
        split_indices = []
        for m in matches:
            t_val = m.group(2)
            start_pos = m.start(2)
            split_indices.append((start_pos, t_val))

        current_time = time_str
        last_end = 0

        for idx, (pos, t_val) in enumerate(split_indices):
            if pos <= 2 and desc[:pos].strip() in ["", "，", ",", ":", "："]:
                current_time = t_val
                last_end = pos + len(t_val)
                continue

            part_desc = desc[last_end:pos].strip(" 。；;,，\n")
            if part_desc:
                sub_events.append((current_time, part_desc))

            current_time = t_val
            last_end = pos + len(t_val)
            while last_end < len(desc) and desc[last_end] in [" ", ":", "：", "、"]:
                last_end += 1

        final_desc = desc[last_end:].strip(" 。；;,，\n")
        if final_desc:
            sub_events.append((current_time, final_desc))

        return sub_events if sub_events else [(time_str, desc)]

    def extract_operations_and_tickets(self) -> Tuple[Dict[str, Any], Dict[str, Any]]:
        """
        提取微观操作流、抢修工作票八元组与初期补水指标（支持跨行粘合、原子时刻解耦拆分与跳闸识别）
        返回: (result_dict, flat_water_metrics_6)
        """
        result: Dict[str, Any] = {
            "events": [],
            "water_replenish": {
                "first_station_1_water_tons": None,
                "first_station_2_water_tons": None,
                "period_last_year_water_tons": None,
                "accumulated_water_tons": None,
                "accumulated_last_year_water_tons": None,
                "substations_filled": None
            }
        }

        flat_water_metrics = {
            "water_sup_bh_s1": None,
            "water_sup_bh_s2": None,
            "water_sup_period_last_year": None,
            "water_sup_accum": None,
            "water_sup_accum_last": None,
            "substations_filled": None
        }

        if not self.table_event_logs:
            return result, flat_water_metrics

        table = self.table_event_logs
        events = []
        current_org = "北海电厂"

        # 1. 表格行智能预处理与跨行碎片向上缝合
        normalized_rows = []
        for row_idx, row in enumerate(table):
            if row_idx == 0:
                continue
            cells = [self.clean_text(c) for c in row]
            if not cells or not any(cells):
                continue

            time_str = cells[0].strip() if len(cells) > 0 else ""
            desc = cells[1].strip() if len(cells) > 1 else ""
            if not desc and not time_str:
                continue

            # 探测组织机构抬头（如 "供热公司："、"北海电厂："）
            org = self.detect_org_header(time_str, desc)
            if org:
                normalized_rows.append((time_str, desc, org))
                continue

            # 跨行残片缝合：如果当前行时间为空，且上一行是未闭合的工作票/工期行
            if not time_str and normalized_rows and not normalized_rows[-1][2]:
                prev_t, prev_d, _ = normalized_rows[-1]
                if (re.search(r"[~～至到-]\s*$", prev_d) or 
                    prev_d.endswith("工期") or 
                    prev_d.endswith("工期:") or 
                    prev_d.endswith("工期：") or
                    (re.match(r"^\d{1,2}/\d{1,2}", desc) and ("工作票" in prev_d or "工期" in prev_d))):
                    normalized_rows[-1] = (prev_t, prev_d + desc, None)
                    continue

            normalized_rows.append((time_str, desc, None))

        # 2. 逐行解析与复合事件拆分
        for time_str, desc, org in normalized_rows:
            if org:
                current_org = org
                continue

            for std_name, aliases in self.KNOWN_ORGS:
                if any(desc.startswith(a + ":") or desc.startswith(a + "：") for a in aliases):
                    current_org = std_name
                    desc = re.sub(r"^[^:：]+[:：]\s*", "", desc)
                    break

            if not desc:
                continue

            combined = (time_str + " " + desc).strip()
            # 补水指标行探测（含换行产生的 "去年同期:22069吨"）
            if "补水量" in combined or "上水换热站" in combined or re.search(r"去年同期\s*[:：]?\s*\d+\s*吨?", combined):
                s1_m = re.search(r"#1\s*首站补水量\s*[:：]\s*(\d+(?:\.\d+)?)\s*吨?", combined)
                if s1_m:
                    val = float(s1_m.group(1))
                    result["water_replenish"]["first_station_1_water_tons"] = val
                    flat_water_metrics["water_sup_bh_s1"] = val

                s2_m = re.search(r"#2\s*首站补水量\s*[:：]\s*(\d+(?:\.\d+)?)\s*吨?", combined)
                if s2_m:
                    val = float(s2_m.group(1))
                    result["water_replenish"]["first_station_2_water_tons"] = val
                    flat_water_metrics["water_sup_bh_s2"] = val

                same_m = re.search(r"(?:去年)?同期补水量\s*[:：]\s*(\d+(?:\.\d+)?)\s*吨?", combined)
                if same_m:
                    val = float(same_m.group(1))
                    result["water_replenish"]["period_last_year_water_tons"] = val
                    flat_water_metrics["water_sup_period_last_year"] = val

                acc_m = re.search(r"累计补水量\s*[:：]?\s*(\d+(?:\.\d+)?)\s*吨?", combined)
                if acc_m:
                    val = float(acc_m.group(1))
                    result["water_replenish"]["accumulated_water_tons"] = val
                    flat_water_metrics["water_sup_accum"] = val

                acc_last_m = re.search(r"去年同期\s*[:：]\s*(\d+(?:\.\d+)?)\s*吨?", combined)
                if acc_last_m:
                    val = float(acc_last_m.group(1))
                    result["water_replenish"]["accumulated_last_year_water_tons"] = val
                    flat_water_metrics["water_sup_accum_last"] = val

                sub_m = re.search(r"完成上水换热站\s*[:：]\s*(\d+)\s*个", combined)
                if sub_m:
                    val = int(sub_m.group(1))
                    result["water_replenish"]["substations_filled"] = val
                    flat_water_metrics["substations_filled"] = val
                continue

            # 智能拆分复合事件为原子事件
            sub_events = self.split_compound_event(time_str, desc)
            for sub_t, sub_d in sub_events:
                is_ticket = "工作票" in sub_d or "签发" in sub_d or "抢修" in sub_d
                if is_ticket:
                    # 智能补全因排版紧凑或字距压缩导致缺失的日期时间空格：DD/MMHH:MM -> DD/MM HH:MM
                    sub_d = re.sub(r"(\d{1,2}/\d{1,2})(\d{1,2}:\d{2})", r"\1 \2", sub_d)
                    sub_d = re.sub(r"(工期)[:：]?\s*(\d)", r"\1 \2", sub_d)

                pipe_size_m = re.search(r"(DN\d+)", sub_d, re.IGNORECASE)
                pipe_size = pipe_size_m.group(1).upper() if pipe_size_m else ""

                # 地点实体抽取：支持复杂多地标、住宅区后缀及动作前缀清洗
                if "特勤大队事故办及港湾隽景" in sub_d:
                    location = "特勤大队事故办及港湾隽景"
                else:
                    loc_m = re.search(r"((?:[\u4e00-\u9fa5]+(?:路|街|巷|道|广场|酒店|大厦|桥|苑|村|小区|港湾隽景|港湾|隽景|湾|景|城|港))(?:\s*交\s*[\u4e00-\u9fa5]+(?:路|街|巷|道))?(?:[东西南北中侧]+)?)", sub_d)
                    location = loc_m.group(1).strip() if loc_m else ""
                    # 剥离动词前缀（开启/关闭/签发）与水网介质前缀（高温水/低温水）
                    location = re.sub(r"^(?:开启|关闭|签发|恢复|开|关|高温水|低温水|蒸汽网)+", "", location).strip()

                time_range_m = re.search(r"工期\s*([0-9/:\s~～-]+)", sub_d)
                time_range = time_range_m.group(1).strip() if time_range_m else ""
                if time_range:
                    # 规范化工期内的时间格式与空格
                    time_range = re.sub(r"(\d{1,2}/\d{1,2})(\d{1,2}:\d{2})", r"\1 \2", time_range)
                    time_range = re.sub(r"\s*([~～])\s*", r" ~ ", time_range)
                    time_range = re.sub(r"\s+", " ", time_range).strip()

                action = "运行操作"
                if "跳闸" in sub_d or "事故" in sub_d or "保护动作" in sub_d:
                    action = "故障跳闸"
                elif "并网" in sub_d or "并入" in sub_d or "并汽" in sub_d:
                    action = "并网并列"
                elif "点火" in sub_d or "点炉" in sub_d:
                    action = "点火"
                elif "抢修" in sub_d or "查漏" in sub_d:
                    action = "查漏抢修"
                elif "开启" in sub_d or sub_d.startswith("开"):
                    action = "开启倒闸"
                elif "关闭" in sub_d or sub_d.startswith("关"):
                    action = "关闭倒闸"
                elif "启动" in sub_d or "投运" in sub_d:
                    action = "启动投运"

                events.append({
                    "time": sub_t,
                    "responsible_org": current_org,
                    "is_ticket": is_ticket,
                    "action": action,
                    "pipe_size": pipe_size,
                    "location": location,
                    "time_range": time_range,
                    "raw_description": sub_d
                })

        result["events"] = events
        return result, flat_water_metrics

    def extract_all(self) -> Dict[str, Any]:
        """运行全套提取器，组装 100+ 工业宽表全物理字段展开大字典"""
        header = self.extract_header()
        biz_date = header.get("biz_date")
        issue_date = header.get("issue_date")
        publish_date = header.get("publish_date")
        duty_staff_str = ",".join(header.get("duty_officers", []))

        weather = self.extract_weather(biz_date=biz_date)
        boiler_turbines = self.extract_boiler_turbine_status()
        first_stations, flat_stations = self.extract_first_stations()
        steam_networks, flat_steams = self.extract_steam_networks()
        coal_stats_dict, flat_coals = self.extract_coal_stats()
        production_stats, network_heat, flat_prods = self.extract_production_and_network_heat()
        ops_and_tickets, flat_waters = self.extract_operations_and_tickets()

        # 宏观四大总量计算
        total_coal = production_stats.get("urban_total_actual_coal")
        if total_coal is None:
            bh_c = production_stats.get("beihai_actual_coal")
            xh_c = production_stats.get("xianghai_actual_coal")
            if bh_c is not None or xh_c is not None:
                total_coal = round((bh_c or 0.0) + (xh_c or 0.0), 2)
            else:
                total_coal = flat_coals.get("coal_main_sum_use") or flat_coals.get("coal_total_use")

        total_power = production_stats.get("urban_total_actual_power")
        if total_power is None:
            bh_p = production_stats.get("beihai_actual_power")
            xh_p = production_stats.get("xianghai_actual_power")
            if bh_p is not None or xh_p is not None:
                total_power = round((bh_p or 0.0) + (xh_p or 0.0), 2)

        total_heat = production_stats.get("urban_total_actual_heat")
        if total_heat is None:
            bh_h = production_stats.get("beihai_actual_heat")
            xh_h = production_stats.get("xianghai_actual_heat")
            if bh_h is not None or xh_h is not None:
                total_heat = round((bh_h or 0.0) + (xh_h or 0.0), 2)

        total_water = flat_waters.get("water_sup_accum")

        # 组装扁平化大宽表结果对象
        doc_data: Dict[str, Any] = {
            # 1. 基础发令
            "project_key": self.project_key,
            "biz_date": biz_date,
            "issue_date": issue_date,
            "publish_date": publish_date,
            "duty_staff": duty_staff_str,
            "duty_officers": duty_staff_str,
            "file_name": self.file_name,

            # 2. 气象全要素展开
            "wth_date": weather.get("forecast_date"),
            "wth_weekday": weather.get("weekday"),
            "wth_condition": weather.get("weather_desc"),
            "wth_temp_min": weather.get("temp_min"),
            "wth_temp_max": weather.get("temp_max"),
            "wth_hum_min": weather.get("humidity_min"),
            "wth_hum_max": weather.get("humidity_max"),
            "wth_wind_dir": weather.get("wind_direction"),
            "wth_wind_scale": weather.get("wind_scale"),
            "wth_raw_text": weather.get("raw_text"),
            "weather": weather,
        }

        # 3. 展开水网首站 36 项物理列
        doc_data.update(flat_stations)
        doc_data["first_stations"] = first_stations

        # 4. 展开蒸汽网口 12 项物理列
        doc_data.update(flat_steams)
        doc_data["steam_networks"] = steam_networks

        # 5. 展开燃煤收耗存 25 项物理列
        doc_data.update(flat_coals)
        doc_data["coal_stats"] = coal_stats_dict.get("coal_items", [])

        # 6 & 7. 展开当日产量与网口热量 26 项物理列
        doc_data.update(flat_prods)
        doc_data["production_stats"] = production_stats
        doc_data["network_heat"] = network_heat

        # 8. 展开初期补水量 6 项物理列
        doc_data.update(flat_waters)
        doc_data["water_replenishment"] = ops_and_tickets.get("water_replenish", {})

        # 9. 宏观四大总量
        doc_data["total_coal_consumed"] = total_coal
        doc_data["total_power_generated"] = total_power
        doc_data["total_heat_supplied"] = total_heat
        doc_data["total_water_replenished"] = total_water

        # 10. 高度动态结构化存根 (设备拓扑与抢修事件日志)
        facilities_list = boiler_turbines.get("facilities", [])
        doc_data["boiler_turbine_status"] = facilities_list
        doc_data["eqp_status_json"] = {
            "matrix": boiler_turbines.get("matrix", {}),
            "flat_map": boiler_turbines.get("flat_map", {}),
            "facilities": facilities_list
        }
        events_list = ops_and_tickets.get("events", [])
        doc_data["operations_and_tickets"] = events_list
        doc_data["log_items_json"] = events_list

        return reconcile_report(doc_data)


def extract_docx_file(file_path: str, project_key: str = "daily_report_25_26") -> Dict[str, Any]:
    """从 DOCX 文件提取结构化调度数据"""
    if docx is None:
        raise RuntimeError("Python 环境中未安装 python-docx 依赖库，请执行: pip install python-docx")
    doc = docx.Document(file_path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for t in doc.tables:
        norm_t = []
        for row in t.rows:
            norm_t.append([c.text.strip() for c in row.cells])
        tables.append(norm_t)

    extractor = UnifiedDispatchExtractor(
        paragraphs=paras,
        tables=tables,
        file_name=os.path.basename(file_path),
        project_key=project_key
    )
    return extractor.extract_all()


def extract_pdf_file(file_path: str, project_key: str = "daily_report_25_26") -> Dict[str, Any]:
    """从 PDF 文件提取结构化调度数据"""
    if pdfplumber is None:
        raise RuntimeError("Python 环境中未安装 pdfplumber 依赖库，请执行: pip install pdfplumber")
    paras = []
    tables = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for line in text.split("\n"):
                line = line.strip()
                if line:
                    paras.append(line)
            extracted_tables = page.extract_tables() or []
            for t in extracted_tables:
                norm_t = [[(cell or "").strip() for cell in row] for row in t]
                tables.append(norm_t)

    extractor = UnifiedDispatchExtractor(
        paragraphs=paras,
        tables=tables,
        file_name=os.path.basename(file_path),
        project_key=project_key
    )
    return extractor.extract_all()


def extract_single_file(file_path: str, project_key: str = "daily_report_25_26") -> Dict[str, Any]:
    """根据文件扩展名自动分发提取器"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        return extract_docx_file(file_path, project_key)
    elif ext == ".pdf":
        return extract_pdf_file(file_path, project_key)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only .docx and .pdf are supported.")
