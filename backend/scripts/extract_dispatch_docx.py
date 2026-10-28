# -*- coding: utf-8 -*-
r"""
主城区调度生产日报 (DOCX) 批量结构化提取工具 (全设备与锅炉拓扑标准化版)
支持桌面图形界面 (GUI) 与命令行 (CLI) 双模运行。

【企业与锅炉设备拓扑标准化规范】：
1. 企业标准口径:
   - 北方热电 (原北方电厂)
   - 金州热电 (原金州电厂)
   - 庄河环海 (原庄河电厂)
   - 金普公司 (原金普热电/金普)
   - 北海电厂、香海电厂、供热公司、调度运行中心
2. 锅炉物理类型与所属单元深度归一:
   - 北海电厂细分为:
     * “北海热电联产” (锅炉全为【汽炉】: #1汽炉 ~ #4汽炉)
     * “北海水炉” (锅炉全为【水炉】: #1水炉 ~ #3水炉)
     * “燃气锅炉” (锅炉全为【燃气炉】: #1燃气炉 ~ #10燃气炉)
   - 香海电厂: 锅炉全为【汽炉】(#1汽炉 ~ #4汽炉)
   - 金州热电: 【#6炉为水炉】(#6水炉)，其余锅炉全为【汽炉】(#1汽炉 ~ #4汽炉)
   - 北方热电: 锅炉全为【汽炉】(#1汽炉 ~ #3汽炉)
   - 庄河环海: 全为【锅炉房锅炉】
"""

import os
import sys
import re
import json
import argparse
import threading
from typing import Dict, Any, List, Optional, Tuple
from datetime import datetime
import docx

try:
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox, scrolledtext
    HAS_TKINTER = True
except ImportError:
    HAS_TKINTER = False


class DispatchDocxExtractor:
    """主城区调度生产日报 DOCX 解析器 (设备与锅炉拓扑标准化版)"""

    KNOWN_ORGS = [
        ("北海电厂", ["北海电厂", "北海", "北海厂", "北海热电", "北海热电联产", "北海水炉", "燃气锅炉"]),
        ("香海电厂", ["香海电厂", "香海", "香海厂", "香海热电"]),
        ("金州热电", ["金州热电", "金州电厂", "金州", "金州厂"]),
        ("北方热电", ["北方热电", "北方电厂", "北方", "北方厂"]),
        ("庄河环海", ["庄河环海", "庄河电厂", "庄河", "庄河厂"]),
        ("金普公司", ["金普公司", "金普热电", "金普", "金普新区"]),
        ("供热公司", ["供热公司", "管网公司", "供热", "管网"]),
        ("大连港", ["大连港"]),
        ("调度运行中心", ["调度中心", "运行中心", "调度运行中心"])
    ]

    CATEGORY_ORG_MAP = {
        "北海": "北海热电联产",
        "水炉": "北海水炉",
        "燃气炉": "燃气锅炉",
        "香海": "香海电厂",
        "金州": "金州热电",
        "北方": "北方热电",
        "庄河": "庄河环海"
    }

    def __init__(self, docx_path: str, project_key: str = "daily_report_25_26"):
        self.docx_path = docx_path
        self.project_key = project_key
        self.doc = docx.Document(docx_path)
        self.paragraphs = [p.text.strip() for p in self.doc.paragraphs if p.text.strip()]
        self.tables = self.doc.tables

    @staticmethod
    def clean_text(text: Optional[str]) -> str:
        """文本清洗：统一全半角、异体字替换、去除多余空白"""
        if not text:
            return ""
        t = text.replace("幷", "并").replace("℃", "℃")
        t = t.replace("～", "~").replace("—", "-").replace("－", "-")
        t = t.replace("：", ":").replace("，", ",").replace("／", "/")
        t = re.sub(r"\s+", " ", t).strip()
        return t

    @staticmethod
    def parse_float(val: Any) -> Optional[float]:
        """安全提取浮点数"""
        if val is None:
            return None
        s = str(val).strip()
        if not s or s in ["-", "--", "/", "null", "None"]:
            return None
        match = re.search(r"[-+]?\d*\.?\d+", s)
        if match:
            try:
                return float(match.group())
            except ValueError:
                return None
        return None

    @staticmethod
    def parse_int(val: Any) -> Optional[int]:
        """安全提取整数"""
        f = DispatchDocxExtractor.parse_float(val)
        return int(f) if f is not None else None

    @staticmethod
    def expand_device_numbers(text: str) -> List[str]:
        """展开设备编号：如 #1-10 或 #2、3"""
        cleaned = DispatchDocxExtractor.clean_text(text)
        if not cleaned:
            return []
        
        results = []
        range_match = re.search(r"#?(\d+)\s*[-~]\s*#?(\d+)", cleaned)
        if range_match:
            start_num = int(range_match.group(1))
            end_num = int(range_match.group(2))
            for n in range(start_num, end_num + 1):
                results.append(f"#{n}")
            return results

        nums = re.findall(r"\d+", cleaned)
        for n in nums:
            results.append(f"#{n}")
        return sorted(list(set(results)), key=lambda x: int(x.replace("#", "")))

    @classmethod
    def standardize_boiler_name(cls, raw_text: str, org_or_unit: str) -> Tuple[str, str, str]:
        """
        根据集团底层拓扑规则精准推导锅炉类型与标准命名：
        返回: (boiler_type, standard_facility_name, sub_unit)
        - 北海电厂：
          * 燃气锅炉 -> 燃气炉, #X燃气炉, 燃气锅炉
          * 北海水炉 -> 水炉, #X水炉, 北海水炉
          * 北海热电联产 -> 汽炉, #X汽炉, 北海热电联产
        - 香海电厂 -> 汽炉, #X汽炉, 香海电厂
        - 金州热电 -> #6炉为水炉(#6水炉)，其余为汽炉(#X汽炉), 金州热电
        - 北方热电 -> 汽炉, #X汽炉, 北方热电
        - 庄河环海 -> 锅炉房锅炉, 庄河锅炉房锅炉, 庄河环海
        """
        t = raw_text.strip()
        num_m = re.search(r"(\d+)", t)
        num = f"#{num_m.group(1)}" if num_m else ""

        # 庄河
        if any(k in org_or_unit for k in ["庄河", "庄河环海", "庄河电厂"]):
            return "锅炉房锅炉", f"庄河{num}锅炉房锅炉" if num else "庄河锅炉房锅炉", "庄河环海"

        # 金州热电：#6炉是水炉，其余全为汽炉
        if any(k in org_or_unit for k in ["金州", "金州热电", "金州电厂"]):
            if num == "#6" or "6" in t or "水炉" in t:
                return "水炉", "#6水炉", "金州热电"
            else:
                return "汽炉", f"{num}汽炉" if num else "金州汽炉", "金州热电"

        # 北方热电：全为汽炉
        if any(k in org_or_unit for k in ["北方", "北方热电", "北方电厂"]):
            return "汽炉", f"{num}汽炉" if num else "北方汽炉", "北方热电"

        # 香海电厂：全为汽炉
        if any(k in org_or_unit for k in ["香海", "香海电厂", "香海厂"]):
            return "汽炉", f"{num}汽炉" if num else "香海汽炉", "香海电厂"

        # 北海电厂：分 北海水炉 / 燃气锅炉 / 北海热电联产汽炉
        if "水炉" in t or "水炉" in org_or_unit:
            return "水炉", f"{num}水炉" if num else "北海水炉", "北海水炉"
        elif "燃气" in t or "燃气" in org_or_unit:
            return "燃气炉", f"{num}燃气炉" if num else "北海燃气炉", "燃气锅炉"
        else:
            return "汽炉", f"{num}汽炉" if num else "北海汽炉", "北海热电联产"

    def extract_meta_and_weather(self) -> Dict[str, Any]:
        """提取元数据与气象数据"""
        data: Dict[str, Any] = {
            "project_key": self.project_key,
            "biz_date": None,
            "issue_date": None,
            "duty_staff": "",
            "wth_date": None,
            "wth_weekday": "",
            "wth_condition": "",
            "wth_temp_min": None,
            "wth_temp_max": None,
            "wth_hum_min": None,
            "wth_hum_max": None,
            "wth_wind_dir": "",
            "wth_wind_scale": "",
            "wth_raw_text": ""
        }

        for p in self.paragraphs[:6]:
            p_clean = self.clean_text(p)
            if "值班人员" in p_clean:
                staff_match = re.search(r"值班人员\s*[:：]\s*(.*?)(?:时间|$)", p_clean)
                if staff_match:
                    raw_staff = staff_match.group(1).strip()
                    staff_list = [s.strip() for s in re.split(r"[\s,、]+", raw_staff) if s.strip()]
                    data["duty_staff"] = ",".join(staff_list)

                date_match = re.search(r"时间\s*[:：]\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", p_clean)
                if date_match:
                    y, m, d = date_match.groups()
                    data["biz_date"] = f"{int(y):04d}-{int(m):02d}-{int(d):02d}"

        for i, p in enumerate(self.paragraphs[:10]):
            p_clean = self.clean_text(p)
            if "天气情况" in p_clean or "相对湿度" in p_clean or "℃" in p_clean:
                target_text = p_clean
                if target_text == "天气情况" and i + 1 < len(self.paragraphs):
                    target_text = self.clean_text(self.paragraphs[i + 1])
                
                data["wth_raw_text"] = target_text

                wth_date_match = re.search(r"(\d{1,2})\s*月\s*(\d{1,2})\s*日", target_text)
                if wth_date_match and data["biz_date"]:
                    wm, wd = wth_date_match.groups()
                    year = int(data["biz_date"].split("-")[0])
                    data["wth_date"] = f"{year:04d}-{int(wm):02d}-{int(wd):02d}"

                weekday_match = re.search(r"(星期[一二三四五六日天])", target_text)
                if weekday_match:
                    data["wth_weekday"] = weekday_match.group(1)

                temp_match = re.search(r"(-?\d+(?:\.\d+)?)\s*℃\s*[-~至到]\s*(-?\d+(?:\.\d+)?)\s*℃", target_text)
                if temp_match:
                    t1, t2 = float(temp_match.group(1)), float(temp_match.group(2))
                    data["wth_temp_min"] = min(t1, t2)
                    data["wth_temp_max"] = max(t1, t2)

                hum_match = re.search(r"相对湿度\s*[:：]?\s*(\d+)\s*%\s*[-~至到]\s*(\d+)\s*%", target_text)
                if hum_match:
                    h1, h2 = int(hum_match.group(1)), int(hum_match.group(2))
                    data["wth_hum_min"] = min(h1, h2)
                    data["wth_hum_max"] = max(h1, h2)

                cond_match = re.search(r"℃\s*([^\d相对湿度]+?)(?:相对湿度|西南风|西北风|东北风|东南风|北风|南风|东风|西风|$)", target_text)
                if cond_match:
                    data["wth_condition"] = cond_match.group(1).strip()

                wind_match = re.search(r"((?:东|西|南|北|偏|微)+风)\s*([\d~级阵风\-]+(?:级)?)", target_text)
                if wind_match:
                    data["wth_wind_dir"] = wind_match.group(1).strip()
                    data["wth_wind_scale"] = wind_match.group(2).strip()

        for p in reversed(self.paragraphs[-6:]):
            p_clean = self.clean_text(p)
            issue_match = re.search(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", p_clean)
            if issue_match:
                y, m, d = issue_match.groups()
                data["issue_date"] = f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
                break

        return data

    def extract_equipment_status(self) -> Dict[str, Any]:
        """
        提取电厂机炉运行状态矩阵（Table 0）
        完全落地锅炉细分类型 (汽炉/水炉/燃气炉) 与单元归属 (北海热电联产/北海水炉/燃气锅炉)
        """
        result = {
            "matrix": {},
            "flat_map": {},
            "devices_detail": {}
        }
        if not self.tables:
            return result

        tbl = self.tables[0]
        headers = [self.clean_text(c.text) for c in tbl.rows[0].cells]
        status_cols = {}
        for col_idx, h in enumerate(headers):
            if "运行" in h:
                status_cols["运行"] = col_idx
            elif "备用" in h:
                status_cols["备用"] = col_idx
            elif "检修" in h:
                status_cols["检修"] = col_idx

        for row in tbl.rows[1:]:
            cells = [self.clean_text(c.text) for c in row.cells]
            if not cells or not cells[0]:
                continue
            raw_cat = cells[0].replace(" ", "")
            std_unit = self.CATEGORY_ORG_MAP.get(raw_cat, raw_cat)

            # 区分机和炉
            if raw_cat in ["北海", "香海", "金州", "北方"]:
                result["matrix"][std_unit] = {
                    "机": {"运行": [], "备用": [], "检修": []},
                    "炉": {"运行": [], "备用": [], "检修": []}
                }
                for status_name, col_idx in status_cols.items():
                    if col_idx < len(cells):
                        val = cells[col_idx]
                        m_part = ""
                        b_part = ""
                        if "机" in val:
                            m_part = val.split("机")[0]
                            rest = val.split("机")[1]
                            if "炉" in rest:
                                b_part = rest.split("炉")[0]
                        elif "炉" in val:
                            b_part = val.split("炉")[0]

                        m_list = self.expand_device_numbers(m_part)
                        b_list = self.expand_device_numbers(b_part)

                        result["matrix"][std_unit]["机"][status_name] = m_list
                        result["matrix"][std_unit]["炉"][status_name] = b_list

                        # 1. 汽机映射
                        for dev in m_list:
                            dev_key = f"{std_unit}_{dev}机"
                            result["flat_map"][dev_key] = status_name
                            result["devices_detail"][dev_key] = {
                                "unit": std_unit, "device_type": "汽机", "status": status_name
                            }

                        # 2. 锅炉标准化细分 (汽炉 vs 水炉)
                        for dev in b_list:
                            b_type, b_std_name, sub_u = self.standardize_boiler_name(dev, raw_cat)
                            dev_key = f"{std_unit}_{b_std_name}"
                            result["flat_map"][dev_key] = status_name
                            result["devices_detail"][dev_key] = {
                                "unit": std_unit, "device_type": b_type, "status": status_name
                            }

            elif raw_cat in ["水炉", "燃气炉"]:
                result["matrix"][std_unit] = {
                    std_unit: {"运行": [], "备用": [], "检修": []}
                }
                for status_name, col_idx in status_cols.items():
                    if col_idx < len(cells):
                        val = cells[col_idx]
                        b_list = self.expand_device_numbers(val)
                        result["matrix"][std_unit][std_unit][status_name] = b_list
                        for dev in b_list:
                            b_type, b_std_name, sub_u = self.standardize_boiler_name(dev, raw_cat)
                            dev_key = f"{std_unit}_{b_std_name}"
                            result["flat_map"][dev_key] = status_name
                            result["devices_detail"][dev_key] = {
                                "unit": std_unit, "device_type": b_type, "status": status_name
                            }

        return result

    def extract_water_stations(self) -> Dict[str, Any]:
        """提取水网首站运行工况（Table 1）"""
        data: Dict[str, Any] = {}
        stations = ["bh_s1", "bh_s2", "bh_wl", "xh_lv1", "xh_lv3", "xh_main"]
        metrics = ["pumps", "flow", "press_sup", "press_ret", "temp_sup", "temp_ret"]
        for st in stations:
            for m in metrics:
                data[f"{st}_{m}"] = None

        if len(self.tables) < 2:
            return data

        tbl = self.tables[1]
        current_plant = ""

        for row in tbl.rows[1:]:
            cells = [self.clean_text(c.text) for c in row.cells]
            if not cells or len(cells) < 4:
                continue

            text_line = " ".join(cells)
            if "北海" in cells[0]:
                current_plant = "bh"
            elif "香海" in cells[0]:
                current_plant = "xh"

            station_key = None
            if "#1首站" in text_line or "#1 首站" in text_line:
                station_key = "bh_s1"
            elif "#2首站" in text_line or "#2 首站" in text_line:
                station_key = "bh_s2"
            elif "水炉" in text_line and current_plant == "bh":
                station_key = "bh_wl"
            elif "#1低真空" in text_line or "#1 低真空" in text_line:
                station_key = "xh_lv1"
            elif "#3低真空" in text_line or "#3 低真空" in text_line:
                station_key = "xh_lv3"
            elif "首站" in text_line and current_plant == "xh":
                station_key = "xh_main"

            if not station_key:
                continue

            mode_str = ""
            flow_val = None
            press_str = ""
            temp_str = ""

            for c in cells:
                if "循环泵" in c:
                    mode_str = c
                elif "/" in c:
                    if press_str == "":
                        press_str = c
                    else:
                        temp_str = c
                elif re.match(r"^\d{3,5}$", c):
                    flow_val = self.parse_float(c)

            if mode_str:
                data[f"{station_key}_pumps"] = self.parse_int(mode_str)
            if flow_val is not None:
                data[f"{station_key}_flow"] = flow_val

            if press_str and "/" in press_str:
                parts = press_str.split("/")
                data[f"{station_key}_press_sup"] = self.parse_float(parts[0])
                data[f"{station_key}_press_ret"] = self.parse_float(parts[1])

            if temp_str and "/" in temp_str:
                parts = temp_str.split("/")
                data[f"{station_key}_temp_sup"] = self.parse_float(parts[0])
                data[f"{station_key}_temp_ret"] = self.parse_float(parts[1])

        return data

    def extract_steam_gateways(self) -> Dict[str, Any]:
        """提取蒸汽网口情况（Table 2）"""
        data: Dict[str, Any] = {}
        gw_keys = ["bh_south", "xb_link", "xh_east", "xh_west"]
        for k in gw_keys:
            data[f"{k}_status"] = None
            data[f"{k}_press"] = None
            data[f"{k}_flow"] = None

        if len(self.tables) < 3:
            return data

        tbl = self.tables[2]
        for row in tbl.rows[1:]:
            cells = [self.clean_text(c.text) for c in row.cells]
            if not cells or len(cells) < 3:
                continue

            name = cells[0].replace(" ", "")
            status = cells[1] if len(cells) > 1 else ""
            press_raw = cells[2] if len(cells) > 2 else ""
            flow_raw = cells[3] if len(cells) > 3 else ""

            if "北海南网" in name:
                data["bh_south_status"] = status or "投运"
                data["bh_south_press"] = self.parse_float(press_raw)
                data["bh_south_flow"] = self.parse_float(flow_raw)
            elif "香北联络" in name:
                data["xb_link_status"] = status or "投运"
                data["xb_link_press"] = self.parse_float(press_raw)
                data["xb_link_flow"] = self.parse_float(flow_raw)
            elif "香海东" in name or "香海东西网" in name:
                data["xh_east_status"] = status or "投运"
                data["xh_west_status"] = status or "投运"

                if "/" in press_raw:
                    pp = press_raw.split("/")
                    data["xh_east_press"] = self.parse_float(pp[0])
                    data["xh_west_press"] = self.parse_float(pp[1])
                else:
                    p = self.parse_float(press_raw)
                    data["xh_east_press"] = p
                    data["xh_west_press"] = p

                if "/" in flow_raw:
                    ff = flow_raw.split("/")
                    data["xh_east_flow"] = self.parse_float(ff[0])
                    data["xh_west_flow"] = self.parse_float(ff[1])
                else:
                    f = self.parse_float(flow_raw)
                    data["xh_east_flow"] = f
                    data["xh_west_flow"] = f

        return data

    def extract_coal_stats(self) -> Dict[str, Any]:
        """提取燃煤统计表（Table 3）"""
        data: Dict[str, Any] = {
            "coal_remark": ""
        }
        entities = ["bh", "xh", "main_sum", "jz", "bf", "zh", "dlg", "total"]
        for e in entities:
            data[f"coal_{e}_in"] = None
            data[f"coal_{e}_use"] = None
            data[f"coal_{e}_stock"] = None

        if len(self.tables) < 4:
            return data

        tbl = self.tables[3]
        name_map = {
            "北海": "bh",
            "香海": "xh",
            "主城区合计": "main_sum",
            "金州": "jz",
            "北方": "bf",
            "庄河": "zh",
            "大连港": "dlg",
            "合计": "total"
        }

        for row in tbl.rows[1:]:
            cells = [self.clean_text(c.text) for c in row.cells]
            if not cells:
                continue

            first_col = cells[0].replace(" ", "")
            if "进煤:" in first_col or "进煤：" in first_col or "汽运车" in first_col:
                if not data["coal_remark"]:
                    data["coal_remark"] = cells[0].strip()
                continue

            for ch_name, e_key in name_map.items():
                if ch_name == first_col:
                    if len(cells) > 1:
                        data[f"coal_{e_key}_in"] = self.parse_float(cells[1])
                    if len(cells) > 2:
                        data[f"coal_{e_key}_use"] = self.parse_float(cells[2])
                    if len(cells) > 3:
                        data[f"coal_{e_key}_stock"] = self.parse_float(cells[3])
                    break

        return data

    def extract_production_outputs(self) -> Dict[str, Any]:
        """提取当日产量与网口热量统计（Table 4）"""
        data: Dict[str, Any] = {
            "prod_bh_steam_coal_plan": None, "prod_bh_steam_coal_act": None,
            "prod_bh_water_coal_plan": None, "prod_bh_water_coal_act": None,
            "prod_bh_power_plan": None, "prod_bh_power_act": None,
            "prod_bh_heat_plan": None, "prod_bh_heat_act": None,
            "prod_xh_steam_coal_plan": None, "prod_xh_steam_coal_act": None,
            "prod_xh_water_coal_plan": None, "prod_xh_water_coal_act": None,
            "prod_xh_power_plan": None, "prod_xh_power_act": None,
            "prod_xh_heat_plan": None, "prod_xh_heat_act": None,
            "prod_main_coal_plan": None, "prod_main_coal_act": None,
            "prod_main_power_plan": None, "prod_main_power_act": None,
            "prod_main_heat_plan": None, "prod_main_heat_act": None,
            "gw_high_temp_water": None,
            "gw_lv1_heat": None,
            "gw_lv3_heat": None,
            "gw_bh_south_steam": None,
            "gw_xh_ew_steam": None,
            "gw_xb_link_steam": None
        }

        if len(self.tables) < 5:
            return data

        for tbl_idx in range(4, min(6, len(self.tables))):
            tbl = self.tables[tbl_idx]
            for row in tbl.rows:
                cells = [self.clean_text(c.text) for c in row.cells]
                if not cells:
                    continue
                first_col = cells[0].replace(" ", "")

                if "汽炉耗煤" in first_col:
                    if len(cells) > 2:
                        data["prod_bh_steam_coal_plan"] = self.parse_float(cells[1])
                        data["prod_bh_steam_coal_act"] = self.parse_float(cells[2])
                    if len(cells) > 4:
                        data["prod_xh_steam_coal_plan"] = self.parse_float(cells[3])
                        data["prod_xh_steam_coal_act"] = self.parse_float(cells[4])
                elif "水炉耗煤" in first_col:
                    if len(cells) > 2:
                        data["prod_bh_water_coal_plan"] = self.parse_float(cells[1])
                        data["prod_bh_water_coal_act"] = self.parse_float(cells[2])
                    if len(cells) > 4:
                        data["prod_xh_water_coal_plan"] = self.parse_float(cells[3])
                        data["prod_xh_water_coal_act"] = self.parse_float(cells[4])
                elif "发电" in first_col and "总量" not in first_col:
                    if len(cells) > 2:
                        data["prod_bh_power_plan"] = self.parse_float(cells[1])
                        data["prod_bh_power_act"] = self.parse_float(cells[2])
                    if len(cells) > 4:
                        data["prod_xh_power_plan"] = self.parse_float(cells[3])
                        data["prod_xh_power_act"] = self.parse_float(cells[4])
                elif "供热量" in first_col and "总量" not in first_col:
                    if len(cells) > 2:
                        data["prod_bh_heat_plan"] = self.parse_float(cells[1])
                        data["prod_bh_heat_act"] = self.parse_float(cells[2])
                    if len(cells) > 4:
                        data["prod_xh_heat_plan"] = self.parse_float(cells[3])
                        data["prod_xh_heat_act"] = self.parse_float(cells[4])

                elif "耗煤总量" in first_col:
                    if len(cells) > 1:
                        data["prod_main_coal_plan"] = self.parse_float(cells[1])
                    if len(cells) > 3:
                        data["prod_main_coal_act"] = self.parse_float(cells[3])
                    elif len(cells) > 2:
                        data["prod_main_coal_act"] = self.parse_float(cells[2])
                elif "发电总量" in first_col:
                    if len(cells) > 1:
                        data["prod_main_power_plan"] = self.parse_float(cells[1])
                    if len(cells) > 3:
                        data["prod_main_power_act"] = self.parse_float(cells[3])
                    elif len(cells) > 2:
                        data["prod_main_power_act"] = self.parse_float(cells[2])
                elif "供热总量" in first_col:
                    if len(cells) > 1:
                        data["prod_main_heat_plan"] = self.parse_float(cells[1])
                    if len(cells) > 3:
                        data["prod_main_heat_act"] = self.parse_float(cells[3])
                    elif len(cells) > 2:
                        data["prod_main_heat_act"] = self.parse_float(cells[2])

                elif "高温水" in first_col:
                    data["gw_high_temp_water"] = self.parse_float(cells[-1])
                elif "#1低真空" in first_col or "#1 低真空" in first_col:
                    data["gw_lv1_heat"] = self.parse_float(cells[-1])
                elif "#3低真空" in first_col or "#3 低真空" in first_col:
                    data["gw_lv3_heat"] = self.parse_float(cells[-1])
                elif "北海南网供汽量" in first_col:
                    data["gw_bh_south_steam"] = self.parse_float(cells[-1])
                elif "香海东西网供汽量" in first_col or "香海东、西网" in first_col:
                    data["gw_xh_ew_steam"] = self.parse_float(cells[-1])
                elif "香北联络" in first_col:
                    data["gw_xb_link_steam"] = self.parse_float(cells[-1])

        return data

    @classmethod
    def detect_org_header(cls, time_col: str, content_col: str) -> Optional[str]:
        """组织分组头识别"""
        c = content_col.strip()
        t = time_col.strip()
        c_clean = c.replace(" ", "").rstrip(":：")
        t_clean = t.replace(" ", "").rstrip(":：")

        if len(c_clean) > 12:
            return None

        for std_name, aliases in cls.KNOWN_ORGS:
            for a in aliases:
                if (c_clean == a or t_clean == a) and not t_clean.replace(":", "").isdigit():
                    return std_name
                if c.endswith((":", "：")) and (c_clean == a or c_clean.startswith(a) or c_clean.endswith(a)):
                    return std_name
        return None

    @classmethod
    def parse_facility_and_action(cls, text: str, org: str) -> Tuple[str, str, str, str, str, str, str, str]:
        """
        全量提取八元组物理实体、锅炉细分类型与动作特征：
        返回: (facility_type, facility_name, action, pipe_spec, location, status_result, boiler_type, sub_unit)
        """
        facility_type = "other"
        facility_name = ""
        action = "运行操作"
        pipe_spec = ""
        location = ""
        status_result = "运行"
        boiler_type = ""
        sub_unit = org

        # 1. 管径规格提取
        spec_m = re.search(r"(DN\d+)", text, re.IGNORECASE)
        if spec_m:
            pipe_spec = spec_m.group(1).upper()

        # 2. 地理位置与十字路口
        loc_m = re.search(r"((?:[\u4e00-\u9fa5]+(?:路|街|巷|道|广场|酒店|大厦|桥|苑|村|小区|港))(?:\s*交\s*[\u4e00-\u9fa5]+(?:路|街|巷|道))?(?:[东西南北中侧]+)?)", text)
        if loc_m:
            location = loc_m.group(1).strip()

        # 3. 物理实体分类与名称
        m_valve = re.search(r"((?:[\u4e00-\u9fa5A-Za-z0-9#]+)?(?:隔离门|放水门|直补门|高减|阀门|放空门|疏水门))", text)
        m_pump = re.search(r"((?:[\u4e00-\u9fa5A-Za-z0-9#]+)?(?:循环泵|循环水泵|水泵|给水泵)[A-Z]?(?:泵)?)", text)
        m_heat_pump = re.search(r"((?:[\u4e00-\u9fa5A-Za-z0-9#]+)?热泵(?:机组)?)", text)
        m_heater = re.search(r"((?:[\u4e00-\u9fa5A-Za-z0-9#]+)?(?:加热器|换热器|凝汽器))", text)
        m_aux = re.search(r"((?:[\u4e00-\u9fa5A-Za-z0-9#]+)?(?:吸风机|引风机|送风机|排粉机|给粉机|变频器))", text)
        m_sub = re.search(r"((?:[\u4e00-\u9fa5]+)?换热站(?:支线)?)", text)
        # 锅炉识别
        m_boiler = re.search(r"(\d+#?炉|#?\d+(?:、\d+)*(?:蒸汽|汽|水|燃气)?炉|水炉|燃气炉)", text)
        # 汽机识别
        m_turbine = re.search(r"(\d+#?机|#?\d+(?:、\d+)*机|汽轮机|汽机)", text)
        # 管网识别
        m_pipe = re.search(r"((?:DN\d+)?(?:主线|支线|东环|西网|东网|南网|北网|蒸汽管道|高温水网|高温水管线|高温水管|管线)[^，,。、\s]*)", text)

        if m_valve:
            facility_type = "valve"
            facility_name = m_valve.group(1).strip()
        elif m_pump:
            facility_type = "pump"
            facility_name = m_pump.group(1).strip()
        elif m_heat_pump:
            facility_type = "heat_pump"
            facility_name = m_heat_pump.group(1).strip()
        elif m_heater:
            facility_type = "heater"
            facility_name = m_heater.group(1).strip()
        elif m_aux:
            facility_type = "auxiliary"
            facility_name = m_aux.group(1).strip()
        elif m_sub:
            facility_type = "substation"
            facility_name = m_sub.group(1).strip()
        elif m_boiler:
            facility_type = "boiler"
            b_raw = m_boiler.group(1).strip()
            # 触发锅炉精准拓扑归一化
            boiler_type, facility_name, sub_unit = cls.standardize_boiler_name(b_raw, org)
        elif m_turbine:
            facility_type = "turbine"
            facility_name = m_turbine.group(1).strip()
            if not facility_name.startswith("#") and re.search(r"\d+", facility_name):
                facility_name = f"#{re.search(r'\d+', facility_name).group()}机"
        elif m_pipe:
            facility_type = "pipeline"
            facility_name = m_pipe.group(1).strip()

        # 供热公司兜底
        if org == "供热公司" and not facility_name and "工作票" in text:
            title_m = re.search(r"[“\"《](.*?)[”\"》]", text)
            if title_m:
                facility_type = "pipeline"
                facility_name = title_m.group(1).replace("高温水", "").replace("查漏抢修", "").replace("泄漏抢修", "").strip()

        # 4. 动作与状态枚举
        if "跳闸" in text:
            action = "事故跳闸"
            status_result = "检修"
        elif "带压堵漏" in text:
            action = "带压堵漏"
            status_result = "抢修中"
        elif "泄漏" in text or "漏泄" in text or "查漏" in text or "管漏" in text:
            action = "查漏抢修"
            status_result = "抢修中"
        elif "抢修" in text:
            action = "抢修消缺"
            status_result = "抢修中"
        elif "消缺" in text:
            action = "检修消缺"
            status_result = "检修中"
        elif "更换" in text:
            action = "管道更换"
            status_result = "施工中"
        elif "上水" in text:
            action = "管网上水"
            status_result = "试运中"
        elif "暖管" in text or "暖机" in text or "倒暖" in text or "送汽" in text:
            action = "暖管送汽"
            status_result = "启动中"
        elif "点火" in text or "点炉" in text:
            action = "点火"
            status_result = "运行"
        elif "并汽" in text or "幷汽" in text:
            action = "并汽"
            status_result = "运行"
        elif "并网" in text or "并入" in text or "并列" in text or "并热网" in text:
            action = "并网并列"
            status_result = "运行"
        elif "冲转" in text:
            action = "冲转"
            status_result = "启动中"
        elif "隔离" in text:
            action = "隔离消缺"
            status_result = "隔离"
        elif "解列" in text or "解汽" in text or "压火" in text or "灭火" in text:
            action = "停运退出"
            status_result = "备用"
        elif "恢复" in text:
            action = "恢复供暖"
            status_result = "恢复运行"
        elif "延期" in text:
            action = "工期延期"
            status_result = "施工延期"
        elif "开启" in text or text.startswith("开") or "开门" in text:
            action = "开启倒闸"
            status_result = "开启"
        elif "关闭" in text or "关门" in text:
            action = "关闭倒闸"
            status_result = "关闭"
        elif "启动" in text or "投运" in text or "投入" in text or text.startswith("投"):
            action = "启动投运"
            status_result = "运行"
        elif "退出" in text or "停运" in text or "停炉" in text or "停机" in text or "打闸" in text:
            action = "停运退出"
            status_result = "备用"
        elif "正常" in text and "运行正常" in text:
            action = "日常巡检"
            status_result = "正常"

        return facility_type, facility_name, action, pipe_spec, location, status_result, boiler_type, sub_unit

    def extract_event_logs_and_metrics(self) -> Dict[str, Any]:
        """提取标准化事件流水与初期指标"""
        result: Dict[str, Any] = {
            "log_items_json": [],
            "log_raw_text": "",
            "water_supply": {
                "water_sup_bh_s1": None,
                "water_sup_bh_s2": None,
                "water_sup_period_last_year": None,
                "water_sup_accum": None,
                "water_sup_accum_last": None,
                "substations_filled": None
            }
        }

        log_tbl = None
        for tbl in self.tables:
            header_text = "".join([c.text for c in tbl.rows[0].cells])
            if "生产运行情况" in header_text or "生产信息记录" in header_text:
                log_tbl = tbl
                break

        if not log_tbl and len(self.tables) >= 5:
            log_tbl = self.tables[-1]

        if not log_tbl:
            return result

        current_org = "北海电厂"
        raw_lines = []
        seq = 1

        for row in log_tbl.rows[1:]:
            cells = [self.clean_text(c.text) for c in row.cells]
            if not cells or not any(cells):
                continue

            time_col = cells[0].strip()
            content_col = cells[1].strip() if len(cells) > 1 else ""

            # 1. 优先探测是否为组织分组头
            new_org = self.detect_org_header(time_col, content_col)
            if new_org:
                current_org = new_org
                continue

            combined = (time_col + " " + content_col).strip()

            # 2. 初期补水量与上水指标提取
            if "补水量" in combined or "上水换热站" in combined:
                raw_lines.append(combined)
                s1_m = re.search(r"#1\s*首站补水量\s*[:：]\s*(\d+(?:\.\d+)?)\s*吨", combined)
                if s1_m:
                    result["water_supply"]["water_sup_bh_s1"] = float(s1_m.group(1))

                s2_m = re.search(r"#2\s*首站补水量\s*[:：]\s*(\d+(?:\.\d+)?)\s*吨", combined)
                if s2_m:
                    result["water_supply"]["water_sup_bh_s2"] = float(s2_m.group(1))

                same_m = re.search(r"同期补水量\s*[:：]\s*(\d+(?:\.\d+)?)\s*吨", combined)
                if same_m:
                    result["water_supply"]["water_sup_period_last_year"] = float(same_m.group(1))

                accum_m = re.search(r"累计补水量\s*[:：]?\s*(\d+(?:\.\d+)?)\s*吨", combined)
                if accum_m:
                    result["water_supply"]["water_sup_accum"] = float(accum_m.group(1))

                accum_last_m = re.search(r"去年同期\s*[:：]?\s*(\d+(?:\.\d+)?)\s*吨", combined)
                if accum_last_m:
                    result["water_supply"]["water_sup_accum_last"] = float(accum_last_m.group(1))

                sub_m = re.search(r"完成上水换热站\s*[:：]?\s*(\d+)\s*个", combined)
                if sub_m:
                    result["water_supply"]["substations_filled"] = int(sub_m.group(1))
                continue

            # 3. 提取实体、动作、锅炉归一化
            f_type, f_name, act, pipe_spec, loc, status_res, b_type, sub_u = self.parse_facility_and_action(combined, current_org)

            # 4. 工作票
            if "工作票" in combined or "抢修" in combined or current_org == "供热公司":
                raw_lines.append(f"[{current_org}] {combined}")
                ticket_item = {
                    "seq": seq,
                    "category": "work_ticket" if ("工作票" in combined or "抢修" in combined) else "operation",
                    "org": current_org,
                    "sub_unit": sub_u,
                    "time": time_col,
                    "ticket_title": "",
                    "facility_type": f_type if f_type != "other" else "pipeline",
                    "facility_name": f_name or "管网设施",
                    "pipe_spec": pipe_spec,
                    "location": loc,
                    "action": act,
                    "status_result": status_res,
                    "time_range": "",
                    "is_delayed": ("延期" in combined),
                    "is_recovered": ("恢复" in combined),
                    "raw_text": combined
                }

                title_m = re.search(r"[“\"《](.*?)[”\"》]", combined)
                if title_m:
                    full_title = title_m.group(1).strip()
                    ticket_item["ticket_title"] = full_title
                    sub_loc = re.search(r"[（(](.*?)[)）]", full_title)
                    if sub_loc:
                        ticket_item["location"] = ticket_item["location"] or sub_loc.group(1)

                period_m = re.search(r"工期\s*[:：]?\s*([0-9/\s:~至到-]+)", combined)
                if period_m:
                    ticket_item["time_range"] = period_m.group(1).strip(" .，,;；")

                result["log_items_json"].append(ticket_item)
                seq += 1
                continue

            # 5. 倒闸操作与设备事件
            if time_col or content_col:
                event_text = content_col if content_col else time_col
                time_val = time_col if content_col else ""
                
                if not time_val:
                    t_match = re.match(r"^(\d{1,2}:\d{2})\s*(.*)", event_text)
                    if t_match:
                        time_val = t_match.group(1)
                        event_text = t_match.group(2)

                raw_line = f"[{current_org} {time_val}] {event_text}".strip()
                raw_lines.append(raw_line)

                evt_category = "incident" if "跳闸" in event_text else "operation"

                sub_events = []
                inline_time_match = re.search(r"[,，;；。\s]*(\d{1,2}:\d{2})\s*(.*)", event_text)
                if inline_time_match:
                    first_part = event_text[:inline_time_match.start()].strip(" ,，;；。")
                    second_time = inline_time_match.group(1)
                    second_part = inline_time_match.group(2).strip(" ,，;；。")
                    if first_part and time_val:
                        sub_events.append({"time": time_val, "action": first_part})
                    elif first_part:
                        sub_events.append({"time": "", "action": first_part})
                    sub_events.append({"time": second_time, "action": second_part})
                else:
                    sub_events.append({
                        "time": time_val,
                        "action": event_text.strip()
                    })

                op_item = {
                    "seq": seq,
                    "category": evt_category,
                    "org": current_org,
                    "sub_unit": sub_u,
                    "time": time_val,
                    "facility_type": f_type,
                    "facility_name": f_name or "主辅机设施",
                    "boiler_type": b_type,
                    "action": act,
                    "status_result": status_res,
                    "pipe_spec": pipe_spec,
                    "location": loc,
                    "content": event_text,
                    "sub_events": sub_events
                }
                result["log_items_json"].append(op_item)
                seq += 1

        result["log_raw_text"] = "\n".join(raw_lines)
        return result

    def extract_all(self) -> Dict[str, Any]:
        """全流程提取并组装为单表标准记录结构"""
        meta_weather = self.extract_meta_and_weather()
        eqp = self.extract_equipment_status()
        water_st = self.extract_water_stations()
        steam_gw = self.extract_steam_gateways()
        coal = self.extract_coal_stats()
        prod = self.extract_production_outputs()
        logs_and_metrics = self.extract_event_logs_and_metrics()

        record: Dict[str, Any] = {
            **meta_weather,
            "eqp_status_json": eqp,
            **water_st,
            **steam_gw,
            **coal,
            **prod,
            **logs_and_metrics["water_supply"],
            "log_items_json": logs_and_metrics["log_items_json"],
            "log_raw_text": logs_and_metrics["log_raw_text"]
        }
        return record


# ======================================================================
# 图形界面 (GUI) 实现
# ======================================================================

class DispatchExtractorGUI:
    """批量提取图形窗口"""

    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("主城区调度生产信息 (DOCX) 批量标准化提取工具 · Phoenix")
        self.root.geometry("860x640")
        self.root.minsize(760, 520)

        self.dir_path_var = tk.StringVar(value="")
        self.output_path_var = tk.StringVar(value=os.path.abspath("dispatch_extracted_batch.json"))
        self.project_key_var = tk.StringVar(value="daily_report_25_26")
        self.is_processing = False

        self._build_ui()

    def _build_ui(self):
        frame_config = ttk.LabelFrame(self.root, text=" 📂 提取配置与路径 ", padding=12)
        frame_config.pack(fill=tk.X, padx=14, pady=10)

        ttk.Label(frame_config, text="DOCX 文件夹:").grid(row=0, column=0, sticky=tk.W, pady=4)
        entry_dir = ttk.Entry(frame_config, textvariable=self.dir_path_var, width=56)
        entry_dir.grid(row=0, column=1, sticky=tk.EW, padx=8, pady=4)
        btn_browse_dir = ttk.Button(frame_config, text="📁 浏览文件夹...", command=self._select_dir)
        btn_browse_dir.grid(row=0, column=2, pady=4)

        ttk.Label(frame_config, text="保存 JSON 结果:").grid(row=1, column=0, sticky=tk.W, pady=4)
        entry_out = ttk.Entry(frame_config, textvariable=self.output_path_var, width=56)
        entry_out.grid(row=1, column=1, sticky=tk.EW, padx=8, pady=4)
        btn_browse_out = ttk.Button(frame_config, text="💾 浏览保存...", command=self._select_output)
        btn_browse_out.grid(row=1, column=2, pady=4)

        frame_action = ttk.Frame(frame_config)
        frame_action.grid(row=2, column=0, columnspan=3, sticky=tk.EW, pady=(8, 0))

        ttk.Label(frame_action, text="项目代号:").pack(side=tk.LEFT)
        ttk.Entry(frame_action, textvariable=self.project_key_var, width=22).pack(side=tk.LEFT, padx=(6, 20))

        self.btn_run = ttk.Button(frame_action, text="🚀 开始批量提取", command=self._start_batch_thread)
        self.btn_run.pack(side=tk.LEFT, padx=6)

        btn_clear = ttk.Button(frame_action, text="🧹 清空日志", command=self._clear_log)
        btn_clear.pack(side=tk.LEFT, padx=6)

        btn_open_folder = ttk.Button(frame_action, text="📂 打开输出目录", command=self._open_output_folder)
        btn_open_folder.pack(side=tk.RIGHT, padx=6)

        frame_config.columnconfigure(1, weight=1)

        frame_status = ttk.Frame(self.root, padding=(14, 0))
        frame_status.pack(fill=tk.X)

        self.lbl_status = ttk.Label(frame_status, text="就绪：请先选择包含 DOCX 文件的文件夹，然后点击【开始批量提取】。")
        self.lbl_status.pack(anchor=tk.W, pady=(0, 4))

        self.progressbar = ttk.Progressbar(frame_status, orient=tk.HORIZONTAL, mode="determinate")
        self.progressbar.pack(fill=tk.X, pady=(0, 8))

        frame_log = ttk.LabelFrame(self.root, text=" 📝 提取详情与日志 (锅炉与设备拓扑标准版) ", padding=8)
        frame_log.pack(fill=tk.BOTH, expand=True, padx=14, pady=(0, 12))

        self.log_text = scrolledtext.ScrolledText(
            frame_log, wrap=tk.WORD, font=("Consolas", 10), bg="#1e1e1e", fg="#d4d4d4", insertbackground="white"
        )
        self.log_text.pack(fill=tk.BOTH, expand=True)

        self.log_text.tag_config("info", foreground="#569cd6")
        self.log_text.tag_config("success", foreground="#4ec9b0")
        self.log_text.tag_config("warn", foreground="#ce9178")
        self.log_text.tag_config("error", foreground="#f44747")
        self.log_text.tag_config("header", foreground="#dcdcaa", font=("Consolas", 10, "bold"))

        self._log("主城区调度生产日报批量提取工具 (锅炉与设备拓扑标准版) 已就绪！\n", "header")
        self._log("规范化口径：\n")
        self._log("  🏢 企业名称对齐: 北方热电、金州热电、庄河环海、金普公司、北海电厂、香海电厂、供热公司。\n")
        self._log("  🔥 锅炉物理分类: 北海热电联产(汽炉)、北海水炉(水炉)、燃气锅炉(燃气炉)；香海/北方全为汽炉；金州#6炉为水炉其余为汽炉；庄河为锅炉房锅炉。\n")
        self._log("  📏 规格无死角抓取: 全量萃取 DN1000/DN500/DN200 等管径规格。\n\n")

    def _select_dir(self):
        d = filedialog.askdirectory(title="选择包含生产信息 docx 文件的文件夹")
        if d:
            self.dir_path_var.set(os.path.normpath(d))
            default_out = os.path.join(d, "dispatch_extracted_batch.json")
            self.output_path_var.set(os.path.normpath(default_out))
            self.lbl_status.config(text=f"已选定文件夹: {d}")

    def _select_output(self):
        f = filedialog.asksaveasfilename(
            title="选择保存 JSON 文件的路径",
            defaultextension=".json",
            filetypes=[("JSON 文件", "*.json"), ("所有文件", "*.*")],
            initialfile="dispatch_extracted_batch.json"
        )
        if f:
            self.output_path_var.set(os.path.normpath(f))

    def _clear_log(self):
        self.log_text.delete("1.0", tk.END)

    def _open_output_folder(self):
        out_path = self.output_path_var.get().strip()
        if out_path and os.path.exists(out_path):
            folder = os.path.dirname(os.path.abspath(out_path))
            os.system(f'explorer "{folder}"')
        elif out_path and os.path.exists(os.path.dirname(out_path)):
            folder = os.path.dirname(os.path.abspath(out_path))
            os.system(f'explorer "{folder}"')
        else:
            messagebox.showinfo("提示", "输出文件尚未生成，请先执行批量提取！")

    def _log(self, text: str, tag: str = ""):
        if tag:
            self.log_text.insert(tk.END, text, tag)
        else:
            self.log_text.insert(tk.END, text)
        self.log_text.see(tk.END)

    def _start_batch_thread(self):
        if self.is_processing:
            return

        folder = self.dir_path_var.get().strip()
        if not folder or not os.path.isdir(folder):
            messagebox.showerror("错误", "请先选择有效的 DOCX 文件夹！")
            return

        out_json = self.output_path_var.get().strip()
        if not out_json:
            messagebox.showerror("错误", "请指定保存 JSON 文件的路径！")
            return

        self.is_processing = True
        self.btn_run.config(state=tk.DISABLED)

        t = threading.Thread(target=self._run_extraction, args=(folder, out_json), daemon=True)
        t.start()

    def _run_extraction(self, folder: str, out_json: str):
        self._log(f"\n🔍 正在扫描文件夹: {folder} ...\n", "info")
        docx_files = []
        for root, _, files in os.walk(folder):
            for f in files:
                if f.endswith(".docx") and not f.startswith("~$"):
                    docx_files.append(os.path.join(root, f))

        total_count = len(docx_files)
        if total_count == 0:
            self.root.after(0, self._on_finish_empty, folder)
            return

        self._log(f"📦 共找到 {total_count} 个 DOCX 文件，开始全自动标准化解析...\n", "info")
        self.root.after(0, lambda: self.progressbar.config(maximum=total_count, value=0))

        results = []
        success_count = 0
        fail_count = 0
        proj_key = self.project_key_var.get().strip() or "daily_report_25_26"

        for idx, fpath in enumerate(docx_files, start=1):
            fname = os.path.basename(fpath)
            self.root.after(0, lambda i=idx, n=fname: self.lbl_status.config(
                text=f"正在解析 [{i}/{total_count}]: {n}"
            ))

            try:
                extractor = DispatchDocxExtractor(fpath, project_key=proj_key)
                record = extractor.extract_all()
                results.append(record)
                success_count += 1

                b_date = record.get("biz_date") or "未知日期"
                staff = record.get("duty_staff") or "未知"
                logs = record.get("log_items_json", [])
                
                org_set = set(l.get("org") for l in logs if l.get("org"))
                org_display = "/".join(list(org_set)[:3]) if org_set else "无事件"

                log_summary = (
                    f"[{idx}/{total_count}] ✅ 成功: {fname}\n"
                    f"   📅 日期: {b_date} | 👷 值班: {staff} | 🏢 责任主体: {org_display} | 📝 标准事件: {len(logs)} 条\n"
                )
                self.root.after(0, self._log, log_summary, "success")

            except Exception as e:
                fail_count += 1
                err_summary = f"[{idx}/{total_count}] ❌ 失败: {fname} -> 错误: {str(e)}\n"
                self.root.after(0, self._log, err_summary, "error")

            self.root.after(0, lambda i=idx: self.progressbar.config(value=i))

        # 保存结果
        try:
            os.makedirs(os.path.dirname(os.path.abspath(out_json)), exist_ok=True)
            with open(out_json, "w", encoding="utf-8") as f:
                json.dump(results, f, ensure_ascii=False, indent=2)
            save_ok = True
        except Exception as e:
            save_ok = False
            self.root.after(0, self._log, f"❌ 保存 JSON 失败: {str(e)}\n", "error")

        self.root.after(0, self._on_finish_all, total_count, success_count, fail_count, out_json, save_ok)

    def _on_finish_empty(self, folder: str):
        self._log(f"⚠️ 在目录 [{folder}] 下未找到任何 .docx 文件！\n", "warn")
        self.lbl_status.config(text="未找到 .docx 文件。")
        self.btn_run.config(state=tk.NORMAL)
        self.is_processing = False
        messagebox.showwarning("提示", f"在所选目录中未找到任何 .docx 文件：\n{folder}")

    def _on_finish_all(self, total: int, success: int, fail: int, out_json: str, save_ok: bool):
        self.btn_run.config(state=tk.NORMAL)
        self.is_processing = False
        self.progressbar.config(value=total)
        self.lbl_status.config(text=f"批量标准化提取完成！成功 {success} / 失败 {fail}。")

        self._log("=" * 65 + "\n", "header")
        self._log(f"🎉 批量提取全部完成！\n", "header")
        self._log(f"   - 总文件数: {total}\n   - 提取成功: {success}\n   - 解析失败: {fail}\n")
        if save_ok:
            self._log(f"💾 标准化数据已保存至:\n   {os.path.abspath(out_json)}\n", "info")
        self._log("=" * 65 + "\n", "header")

        if save_ok:
            msg = f"批量提取完成！\n\n共处理 {total} 个文件\n成功: {success} 个\n失败: {fail} 个\n\n数据已保存至:\n{out_json}"
            if messagebox.askyesno("提取完成", msg + "\n\n是否立即打开保存目录？"):
                self._open_output_folder()
        else:
            messagebox.showerror("错误", f"提取完成，但在写入 JSON 文件时发生错误，请查看日志！")


def launch_gui():
    """启动图形界面"""
    if not HAS_TKINTER:
        print("❌ 当前环境未检测到 tkinter，无法启动图形界面，请使用命令行参数 --dir 或 --file。")
        sys.exit(1)
    root = tk.Tk()
    try:
        from ctypes import windll
        windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    app = DispatchExtractorGUI(root)
    root.mainloop()


def run_cli(args):
    """纯命令行运行模式"""
    files_to_process = []
    if args.file:
        if os.path.exists(args.file):
            files_to_process.append(args.file)
        else:
            print(f"❌ 找不到指定文件: {args.file}")
            sys.exit(1)
    elif args.dir:
        if os.path.isdir(args.dir):
            for root, _, files in os.walk(args.dir):
                for f in files:
                    if f.endswith(".docx") and not f.startswith("~$"):
                        files_to_process.append(os.path.join(root, f))
        else:
            print(f"❌ 找不到指定目录: {args.dir}")
            sys.exit(1)

    all_records = []
    print(f"🚀 [CLI 模式] 开始标准化处理，共找到 {len(files_to_process)} 个文件...")

    for idx, fpath in enumerate(files_to_process, 1):
        try:
            extractor = DispatchDocxExtractor(fpath, project_key=args.project_key)
            record = extractor.extract_all()
            all_records.append(record)
            print(f"[{idx}/{len(files_to_process)}] ✅ {os.path.basename(fpath)} -> 日期: {record.get('biz_date')}")
        except Exception as e:
            print(f"[{idx}/{len(files_to_process)}] ❌ {os.path.basename(fpath)} -> 失败: {e}")

    output_path = args.output
    with open(output_path, "w", encoding="utf-8") as out_f:
        json.dump(all_records, out_f, ensure_ascii=False, indent=2)

    print(f"\n🎉 提取完成！共提取 {len(all_records)} 条标准化数据，已保存至: {os.path.abspath(output_path)}")


def main():
    parser = argparse.ArgumentParser(description="主城区调度生产日报 DOCX 批量标准化提取工具")
    parser.add_argument("--file", "-f", type=str, help="单个 docx 文件路径")
    parser.add_argument("--dir", "-d", type=str, help="包含 docx 文件的目录路径")
    parser.add_argument("--output", "-o", type=str, default="dispatch_extracted_batch.json", help="输出 JSON 路径")
    parser.add_argument("--project-key", "-p", type=str, default="daily_report_25_26", help="项目代号")
    parser.add_argument("--gui", action="store_true", help="强制启动图形界面")

    if len(sys.argv) == 1:
        launch_gui()
        return

    args = parser.parse_args()
    if args.gui or (not args.file and not args.dir):
        launch_gui()
    else:
        run_cli(args)


if __name__ == "__main__":
    main()
