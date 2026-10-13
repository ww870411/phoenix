"""
辽宁电力现货15分钟电价 Excel 数据导入脚本。

功能：
1. 读取 configs/10.13 work_辽宁电力现货15分钟电价_2026.xlsx 的“15分钟电价”工作表；
2. 解析交易日期、时点序号(1-96)、官方时刻、日前电价、实时电价；
3. 自动计算时段起止时间 (period_start_time, period_end_time) 与现货价差 (price_spread)；
4. 采用 PostgreSQL 的 ON CONFLICT (trade_date, point_index) DO UPDATE 机制进行幂等批量写入；
5. 若同一时点已存在售电量填报，更新电价时保留原机组售电量，保证业务数据安全。
"""

from __future__ import annotations

import argparse
import sys
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from sqlalchemy import text

# 将项目根目录加入 sys.path，支持独立脚本执行
_project_root = Path(__file__).resolve().parents[2]
if str(_project_root) not in sys.path:
    sys.path.insert(0, str(_project_root))

from backend.db.database_daily_report_25_26 import engine


def _calculate_period_times(point_index: int) -> tuple[str, str]:
    """
    根据时点序号 (1-96) 计算 15 分钟区间的起止时刻字符串。
    1 -> ('00:00:00', '00:15:00')
    96 -> ('23:45:00', '24:00:00')
    """
    total_minutes_start = (point_index - 1) * 15
    start_hour = total_minutes_start // 60
    start_minute = total_minutes_start % 60
    start_str = f"{start_hour:02d}:{start_minute:02d}:00"

    total_minutes_end = point_index * 15
    if total_minutes_end >= 24 * 60:
        end_str = "24:00:00"
    else:
        end_hour = total_minutes_end // 60
        end_minute = total_minutes_end % 60
        end_str = f"{end_hour:02d}:{end_minute:02d}:00"

    return start_str, end_str


def parse_excel_rows(excel_path: Path, sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
    """解析 Excel 文件中的 15 分钟电价数据。"""
    if not excel_path.exists():
        raise FileNotFoundError(f"未找到指定的 Excel 文件：{excel_path}")

    wb = openpyxl.load_workbook(str(excel_path), read_only=True, data_only=True)
    if sheet_name and sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
    else:
        sheet = wb.active

    records: List[Dict[str, Any]] = []
    header_found = False
    col_map: Dict[str, int] = {}

    for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
        if row_idx == 0 or not header_found:
            # 解析表头定位各列下标
            row_str = [str(cell).strip() if cell is not None else "" for cell in row]
            if "交易日期" in row_str and "时点序号" in row_str:
                for idx, name in enumerate(row_str):
                    if "交易日期" in name:
                        col_map["trade_date"] = idx
                    elif "时点序号" in name:
                        col_map["point_index"] = idx
                    elif "官方时刻" in name:
                        col_map["time_label"] = idx
                    elif "日前电价" in name:
                        col_map["day_ahead_price"] = idx
                    elif "实时电价" in name:
                        col_map["real_time_price"] = idx
                header_found = True
                continue

        # 数据行解析
        raw_date = row[col_map["trade_date"]] if "trade_date" in col_map else None
        raw_point = row[col_map["point_index"]] if "point_index" in col_map else None
        raw_time = row[col_map["time_label"]] if "time_label" in col_map else None
        raw_da = row[col_map["day_ahead_price"]] if "day_ahead_price" in col_map else None
        raw_rt = row[col_map["real_time_price"]] if "real_time_price" in col_map else None

        if raw_date is None or raw_point is None:
            continue

        # 解析日期
        if isinstance(raw_date, datetime):
            parsed_date = raw_date.date()
        elif isinstance(raw_date, date):
            parsed_date = raw_date
        else:
            parsed_date = datetime.strptime(str(raw_date).strip()[:10], "%Y-%m-%d").date()

        point_index = int(raw_point)
        time_label = str(raw_time).strip() if raw_time is not None else f"{point_index}"

        # 起止时间
        p_start, p_end = _calculate_period_times(point_index)

        # 电价转换 (Decimal 避免浮点误差)
        da_price = Decimal(str(raw_da)) if raw_da is not None and str(raw_da).strip() != "" else None
        rt_price = Decimal(str(raw_rt)) if raw_rt is not None and str(raw_rt).strip() != "" else None
        spread = (rt_price - da_price) if (rt_price is not None and da_price is not None) else None

        records.append({
            "trade_date": parsed_date,
            "point_index": point_index,
            "time_label": time_label,
            "period_start_time": p_start,
            "period_end_time": p_end,
            "day_ahead_price": da_price,
            "real_time_price": rt_price,
            "price_spread": spread,
        })

    wb.close()
    return records


UPSERT_SQL = text("""
INSERT INTO public.spot_electricity_15min_data (
    trade_date,
    point_index,
    time_label,
    period_start_time,
    period_end_time,
    day_ahead_price,
    real_time_price,
    price_spread,
    status,
    updated_at
) VALUES (
    :trade_date,
    :point_index,
    :time_label,
    :period_start_time,
    :period_end_time,
    :day_ahead_price,
    :real_time_price,
    :price_spread,
    'raw',
    CURRENT_TIMESTAMP
)
ON CONFLICT (trade_date, point_index)
DO UPDATE SET
    time_label = EXCLUDED.time_label,
    period_start_time = EXCLUDED.period_start_time,
    period_end_time = EXCLUDED.period_end_time,
    day_ahead_price = EXCLUDED.day_ahead_price,
    real_time_price = EXCLUDED.real_time_price,
    price_spread = EXCLUDED.price_spread,
    updated_at = CURRENT_TIMESTAMP;
""")


def import_records_to_db(records: List[Dict[str, Any]], batch_size: int = 1000) -> int:
    """批量幂等导入数据至数据库。"""
    total = len(records)
    if total == 0:
        return 0

    inserted_count = 0
    with engine.begin() as conn:
        for i in range(0, total, batch_size):
            batch = records[i : i + batch_size]
            conn.execute(UPSERT_SQL, batch)
            inserted_count += len(batch)

    return inserted_count


def run_import(excel_path: Path, batch_size: int = 1000) -> Dict[str, Any]:
    """主执行函数，返回统计信息。"""
    print(f"[*] 正在读取 Excel: {excel_path} ...")
    records = parse_excel_rows(excel_path, sheet_name="15分钟电价")
    print(f"[*] 成功解析 {len(records)} 条时序记录。")

    if not records:
        print("[!] 未解析到任何记录，终止导入。")
        return {"total_records": 0, "imported_count": 0}

    dates = sorted({r["trade_date"] for r in records})
    start_date = dates[0]
    end_date = dates[-1]

    da_values = [r["day_ahead_price"] for r in records if r["day_ahead_price"] is not None]
    rt_values = [r["real_time_price"] for r in records if r["real_time_price"] is not None]

    avg_da = (sum(da_values) / len(da_values)) if da_values else Decimal("0")
    avg_rt = (sum(rt_values) / len(rt_values)) if rt_values else Decimal("0")

    print(f"[*] 数据覆盖范围: {start_date} 至 {end_date} (共 {len(dates)} 天，预期每天 96 点)")
    print(f"[*] 日前电价均值: {avg_da:.2f} 元/兆瓦时, 实时电价均值: {avg_rt:.2f} 元/兆瓦时")
    print(f"[*] 开始批量导入 PostgreSQL (批大小: {batch_size}) ...")

    imported_count = import_records_to_db(records, batch_size=batch_size)
    print(f"[✓] 成功导入/更新 {imported_count} 条记录至表 public.spot_electricity_15min_data！")

    return {
        "excel_path": str(excel_path),
        "total_records": len(records),
        "imported_count": imported_count,
        "distinct_days": len(dates),
        "start_date": str(start_date),
        "end_date": str(end_date),
        "avg_day_ahead_price": float(avg_da),
        "avg_real_time_price": float(avg_rt),
    }


def main():
    parser = argparse.ArgumentParser(description="导入辽宁电力现货15分钟电价数据至 PostgreSQL")
    parser.add_argument(
        "--excel-path",
        type=Path,
        default=_project_root / "configs" / "10.13 work_辽宁电力现货15分钟电价_2026.xlsx",
        help="Excel 文件路径",
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=1000,
        help="分批写入批次大小",
    )
    args = parser.parse_args()

    try:
        run_import(args.excel_path, batch_size=args.batch_size)
    except Exception as exc:
        print(f"[✗] 导入发生异常: {exc}", file=sys.stderr)
        raise exc


if __name__ == "__main__":
    main()
