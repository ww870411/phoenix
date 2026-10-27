"""导入 25-26 供暖期线性模型优化后建议 k 参数到 public.dynamic_month_plan_k 数据表。

数据源: configs/10.27（4） 完整指标_最新情况导入用.xlsx
目标表: public.dynamic_month_plan_k
"""

import os
import sys
from typing import Dict, List, Tuple
import openpyxl
import psycopg2
from psycopg2.extras import execute_batch

# 统计口径中英映射
COMPANY_MAP: Dict[str, str] = {
    "集团全口径": "Group",
    "主城区": "ZhuChengQu",
    "北海热电厂": "BeiHai",
    "香海热电厂": "XiangHai",
    "供热公司": "GongRe",
    "金州热电": "JinZhou",
    "北方热电": "BeiFang",
    "金普热电": "JinPu",
    "庄河环海": "ZhuangHe",
    "研究院": "YanJiuYuan",
}

# 生产指标中英映射
ITEM_MAP: Dict[str, str] = {
    "供热量": "amount_heat_supply",
    "发电量": "amount_power_generation",
    "售电量": "amount_power_sales",
    "售汽量": "amount_steam_sales",
    "高温水销售量": "amount_hot_water_sales",
    "原煤耗量": "consumption_amount_raw_coal",
    "标煤耗量": "consumption_std_coal",
    "外购热量": "amount_purchased_heat",
    "站内耗热量": "consumption_station_heat",
    "供暖热单耗": "rate_heat_per_10k_m2",
    "供暖电单耗": "rate_power_per_10k_m2",
}

def load_data_from_excel(excel_path: str) -> List[Tuple]:
    """从 Excel 中加载并校验 360 条模型参数记录。"""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"未找到 Excel 文件: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["完整指标与优化记录"]

    records = []
    for r in range(6, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 7)]
        if vals[0] is None:
            continue

        model_no = int(vals[0])
        comp_cn = str(vals[1]).strip()
        month = int(vals[2])
        item_cn = str(vals[3]).strip()
        unit = str(vals[4]).strip() if vals[4] is not None else ""
        k_val = float(vals[5]) if vals[5] is not None else 0.0

        comp_en = COMPANY_MAP.get(comp_cn)
        item_en = ITEM_MAP.get(item_cn)

        if not comp_en:
            raise ValueError(f"第 {r} 行遇到未知统计口径: {comp_cn}")
        if not item_en:
            raise ValueError(f"第 {r} 行遇到未知指标名称: {item_cn}")

        records.append((
            model_no,
            comp_en,
            comp_cn,
            month,
            item_en,
            item_cn,
            unit,
            k_val,
        ))

    return records


def import_to_postgres(records: List[Tuple], db_url: str):
    """将参数记录批量幂等写入 PostgreSQL 数据库。"""
    upsert_sql = """
    INSERT INTO public.dynamic_month_plan_k (
        model_no, company, company_cn, month, item, item_cn, unit, k_value
    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
    ON CONFLICT (company, month, item) DO UPDATE SET
        model_no = EXCLUDED.model_no,
        company_cn = EXCLUDED.company_cn,
        item_cn = EXCLUDED.item_cn,
        unit = EXCLUDED.unit,
        k_value = EXCLUDED.k_value,
        created_at = CURRENT_TIMESTAMP;
    """

    conn = psycopg2.connect(db_url)
    try:
        with conn.cursor() as cur:
            execute_batch(cur, upsert_sql, records, page_size=100)
        conn.commit()
        print(f"[*] 成功将 {len(records)} 条模型参数写入/更新到 public.dynamic_month_plan_k 表！")
    finally:
        conn.close()


def main():
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    excel_path = os.path.join(base_dir, "configs", "10.27（4） 完整指标_最新情况导入用.xlsx")
    db_url = os.getenv("DATABASE_URL", "postgresql://postgres:postgres@localhost:5432/phoenix")

    print(f"[*] 正在从 {excel_path} 读取模型参数...")
    records = load_data_from_excel(excel_path)
    print(f"[+] 校验完成，读取到 {len(records)} 条模型记录。")

    print("[*] 正在写入 PostgreSQL 数据库...")
    import_to_postgres(records, db_url)
    print("[✓] 全部数据入库成功！")


if __name__ == "__main__":
    main()
