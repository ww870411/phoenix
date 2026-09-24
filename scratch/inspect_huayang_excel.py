# -*- coding: utf-8 -*-
import os
import sys
import openpyxl

file_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "configs", "9.23_导入_辽宁华阳散货.xlsx")

def inspect_excel():
    print(f"检查文件: {file_path}")
    if not os.path.exists(file_path):
        print("文件不存在！")
        return
    wb = openpyxl.load_workbook(file_path, data_only=True)
    print(f"工作表列表: {wb.sheetnames}")
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n--- Sheet: {name} (共 {ws.max_row} 行, {ws.max_column} 列) ---")
        for r in range(1, min(ws.max_row + 1, 30)):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            # 过滤掉末尾全 None
            while row_vals and row_vals[-1] is None:
                row_vals.pop()
            if any(v is not None for v in row_vals):
                print(f"Row {r}: {row_vals}")

if __name__ == "__main__":
    inspect_excel()
