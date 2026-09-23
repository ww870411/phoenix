# -*- coding: utf-8 -*-
import os
import openpyxl

file_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "configs", "9.22_导入_天津天地龙管业.xlsx"))
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb["标准化价格表"]

headers = [cell for cell in next(ws.iter_rows(values_only=True))]
print(f"Headers: {headers}")

target_rows = []
for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if idx == 1:
        continue
    spec = str(row[3] or "")
    if "（甲供钢管）" in spec or "甲供钢管" in spec:
        target_rows.append((idx, row))

print(f"\n找到符合条件的行数: {len(target_rows)}")
for idx, r in target_rows:
    print(f"Row {idx}:")
    for h, v in zip(headers, r):
        print(f"  {h}: {v}")
