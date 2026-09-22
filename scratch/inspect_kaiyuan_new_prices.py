# -*- coding: utf-8 -*-
"""
探查《configs/9.22 导入_开元新增高温水3、4标段保温管、管件.xlsx》的结构与内容。
"""
import os
import sys
import openpyxl

project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if project_root not in sys.path:
    sys.path.insert(0, project_root)

excel_path = r"D:\编程项目\phoenix\configs\9.22 导入_开元新增高温水3、4标段保温管、管件.xlsx"

def inspect():
    if not os.path.exists(excel_path):
        print(f"❌ 文件不存在: {excel_path}")
        return

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    print(f"工作表列表: {wb.sheetnames}")

    ws = wb["产品明细"]
    print(f"总行数: {ws.max_row}")

    categories = set()
    mat_names = set()
    units = set()
    suppliers = set()
    pipe_rows = []
    fitting_rows = []

    for r in range(2, ws.max_row + 1):
        seq = ws.cell(r, 1).value
        sup = ws.cell(r, 2).value
        cat = ws.cell(r, 3).value
        mat = ws.cell(r, 4).value
        spec = ws.cell(r, 5).value
        unit = ws.cell(r, 6).value
        qty = ws.cell(r, 7).value
        price = ws.cell(r, 8).value
        rem = ws.cell(r, 9).value

        if not sup and not spec and not price:
            continue

        sup_str = str(sup or "").strip()
        cat_str = str(cat or "").strip()
        mat_str = str(mat or "").strip()
        spec_str = str(spec or "").strip()
        unit_str = str(unit or "").strip()

        suppliers.add(sup_str)
        categories.add(cat_str)
        mat_names.add(mat_str)
        units.add(unit_str)

        kind = "pipe" if cat_str == "保温管" or unit_str == "米" else "fitting"
        item = {
            "row": r,
            "seq": seq,
            "sup": sup_str,
            "cat": cat_str,
            "mat": mat_str,
            "spec": spec_str,
            "unit": unit_str,
            "qty": qty,
            "price": price,
            "rem": rem,
            "kind": kind
        }
        if kind == "pipe":
            pipe_rows.append(item)
        else:
            fitting_rows.append(item)

    print(f"\n有效数据总行数: {len(pipe_rows) + len(fitting_rows)}")
    print(f"  - 保温管 (pipe) 条数: {len(pipe_rows)}")
    print(f"  - 管件 (fitting) 条数: {len(fitting_rows)}")
    print(f"涉及供应商: {suppliers}")
    print(f"涉及物理类别: {categories}")
    print(f"涉及计量单位: {units}")

    from backend.db.database_daily_report_25_26 import SessionLocal
    from sqlalchemy import text
    session = SessionLocal()
    try:
        existing_kaiyuan = session.execute(text("""
            SELECT material_kind, category, material_name, model_spec, raw_model_spec, unit, unit_price
            FROM tube.tube_material_price
            WHERE supplier_name ILIKE '%开元%'
            ORDER BY material_kind DESC, category, id
            LIMIT 15;
        """)).fetchall()
        print("\n=== 库中现有开元单价样本 (高温水1、2标段) ===")
        for ek in existing_kaiyuan:
            print(f"  kind={ek[0]} | cat={ek[1]} | name={ek[2]} | model_spec={ek[3]} | raw_spec={ek[4]} | unit={ek[5]} | price={ek[6]}")
    finally:
        session.close()

if __name__ == "__main__":
    inspect()
