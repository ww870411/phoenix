# -*- coding: utf-8 -*-
"""
查看 configs/9.28 导入_河北泽越球阀采购价格表.xlsx 的表格结构与数据明细。
"""
import os
import openpyxl

EXCEL_PATH = r"D:\编程项目\phoenix\configs\9.28 导入_河北泽越球阀采购价格表.xlsx"

def inspect():
    if not os.path.exists(EXCEL_PATH):
        print(f"文件不存在: {EXCEL_PATH}")
        return

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    print("工作表列表 (Sheets):", wb.sheetnames)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n==================== Sheet: {sheet_name} ====================")
        print(f"最大行数: {ws.max_row}, 最大列数: {ws.max_column}")

        # 打印全部行数据
        print("\n--- 全部行数据 ---")
        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            print(f"Row {r:2d}: {row_vals}")

        # 检查是否有合并单元格
        if ws.merged_cells.ranges:
            print(f"\n合并单元格范围 ({len(ws.merged_cells.ranges)} 处):")
            for m in list(ws.merged_cells.ranges)[:15]:
                print(f"  - {m.coord}")

if __name__ == "__main__":
    inspect()
