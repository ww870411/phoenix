# -*- coding: utf-8 -*-
import os
import openpyxl

file_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "configs", "9.22_导入_天津天地龙管业.xlsx"))
print(f"Checking Excel file: {file_path}")
if not os.path.exists(file_path):
    print("File does not exist!")
    exit(1)

wb = openpyxl.load_workbook(file_path, data_only=True)
print(f"Sheet names: {wb.sheetnames}")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} (rows: {ws.max_row}, cols: {ws.max_column}) ---")
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    # 打印前5行看看表头
    for idx, r in enumerate(rows[:5]):
        print(f"Row {idx+1}: {r}")
    
    # 查找带有“甲供钢管”的行
    matches = []
    for r_idx, r in enumerate(rows):
        for c_idx, cell in enumerate(r):
            if cell and "甲供钢管" in str(cell):
                matches.append((r_idx + 1, c_idx + 1, r))
                break
    print(f"Total rows containing '甲供钢管': {len(matches)}")
    for r_idx, c_idx, r in matches[:20]:
        print(f"  Row {r_idx}: {r}")
